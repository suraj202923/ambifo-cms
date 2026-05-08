import csv
import io
from pathlib import Path

from app.models import (
    GatheringBlockStorageDetail,
    GatheringFileNasDetail,
    GatheringServerDetail,
    db,
)


SERVER_FIELD_MAP = {
    "server name": "server_name",
    "cpu cores": "cpu_cores",
    "memory (mb)": "memory_mb",
    "provisioned storage (gb)": "provisioned_storage_gb",
    "operating system": "operating_system",
    "is virtual?": "is_virtual",
    "is virtual (yes/no)": "is_virtual",
    "hypervisor name": "hypervisor_name",
    "cpu string": "cpu_string",
    "environment": "environment",
    "sql edition": "sql_edition",
    "application": "application",
    "cpu utilization peak (%)": "cpu_utilization_peak",
    "cpu util peak (%)": "cpu_utilization_peak",
    "memory utilization peak (%)": "memory_utilization_peak",
    "memory util peak (%)": "memory_utilization_peak",
    "time in-use (%)": "time_in_use",
    "annual cost (usd)": "annual_cost_usd",
    "storage type": "storage_type",
}

FILE_NAS_FIELD_MAP = {
    "file server/share name": "file_server_share_name",
    "total used capacity (gb) - usable": "total_used_capacity_gb",
    "access protocol (cifs/nfs)": "access_protocol",
    "total provisioned capacity (gb) - usable": "total_provisioned_capacity_gb",
    "storage efficiency ratio (dedupe/compression)": "storage_efficiency_ratio",
    "peak  iops (reads & writes)": "peak_iops",
    "peak iops (reads & writes)": "peak_iops",
    "peak throughput (mbps)": "peak_throughput_mbps",
    "average  iops (reads & writes)": "average_iops",
    "average iops (reads & writes)": "average_iops",
    "average throughput (mbps)": "average_throughput_mbps",
    "storage pool name": "storage_pool_name",
    "array name": "array_name",
    "array vendor": "array_vendor",
    "average latency (ms)": "average_latency_ms",
    "application": "application",
}

BLOCK_STORAGE_FIELD_MAP = {
    "volume name": "volume_name",
    "total used capacity (gb) - usable": "total_used_capacity_gb",
    "total provisioned capacity (gb) - usable": "total_provisioned_capacity_gb",
    "peak  iops (reads & writes)": "peak_iops",
    "peak iops (reads & writes)": "peak_iops",
    "peak throughput (mbps)": "peak_throughput_mbps",
    "average  iops (reads & writes)": "average_iops",
    "average iops (reads & writes)": "average_iops",
    "average throughput (mbps)": "average_throughput_mbps",
    "array name": "array_name",
    "average latency (ms)": "average_latency_ms",
    "application": "application",
}


def _to_float_safe(value):
    raw = str(value or "").strip()
    if not raw:
        return None
    return float(raw)


def _to_int_safe(value):
    raw = str(value or "").strip()
    if not raw:
        return None
    return int(float(raw))


def _normalize_row(row):
    mapped = {}
    for key, value in (row or {}).items():
        k = (str(key or "")).strip().lower()
        mapped[k] = str(value or "").strip()
    return mapped


def _import_server_rows(rows_of_dicts, customer_id, gathering_request_id=None):
    added = 0
    errors = []
    for i, row in enumerate(rows_of_dicts):
        norm = _normalize_row(row)
        mapped = {}
        for raw_col, val in norm.items():
            field = SERVER_FIELD_MAP.get(raw_col)
            if field:
                mapped[field] = val

        sname = mapped.get("server_name", "")
        if not sname:
            continue
        try:
            db.session.add(GatheringServerDetail(
                customer_id=customer_id,
                gathering_request_id=gathering_request_id,
                server_name=sname,
                cpu_cores=_to_int_safe(mapped.get("cpu_cores")),
                memory_mb=_to_int_safe(mapped.get("memory_mb")),
                provisioned_storage_gb=_to_float_safe(mapped.get("provisioned_storage_gb")),
                operating_system=mapped.get("operating_system") or None,
                is_virtual=(mapped.get("is_virtual") or "").lower() in ("yes", "true", "1"),
                hypervisor_name=mapped.get("hypervisor_name") or None,
                cpu_string=mapped.get("cpu_string") or None,
                environment=mapped.get("environment") or None,
                sql_edition=mapped.get("sql_edition") or None,
                application=mapped.get("application") or None,
                storage_type=mapped.get("storage_type") or None,
                cpu_utilization_peak=_to_float_safe(mapped.get("cpu_utilization_peak")),
                memory_utilization_peak=_to_float_safe(mapped.get("memory_utilization_peak")),
                time_in_use=_to_float_safe(mapped.get("time_in_use")),
                annual_cost_usd=_to_float_safe(mapped.get("annual_cost_usd")),
            ))
            added += 1
        except ValueError as exc:
            errors.append(f"Server row {i + 2}: {exc} - skipped.")
    return added, errors


def _import_file_nas_rows(rows_of_dicts, customer_id, gathering_request_id=None):
    added = 0
    errors = []
    for i, row in enumerate(rows_of_dicts):
        norm = _normalize_row(row)
        mapped = {}
        for raw_col, val in norm.items():
            field = FILE_NAS_FIELD_MAP.get(raw_col)
            if field:
                mapped[field] = val

        name = mapped.get("file_server_share_name", "")
        if not name:
            continue
        try:
            db.session.add(GatheringFileNasDetail(
                customer_id=customer_id,
                gathering_request_id=gathering_request_id,
                file_server_share_name=name,
                total_used_capacity_gb=_to_float_safe(mapped.get("total_used_capacity_gb")),
                access_protocol=mapped.get("access_protocol") or None,
                total_provisioned_capacity_gb=_to_float_safe(mapped.get("total_provisioned_capacity_gb")),
                storage_efficiency_ratio=_to_float_safe(mapped.get("storage_efficiency_ratio")),
                peak_iops=_to_float_safe(mapped.get("peak_iops")),
                peak_throughput_mbps=_to_float_safe(mapped.get("peak_throughput_mbps")),
                average_iops=_to_float_safe(mapped.get("average_iops")),
                average_throughput_mbps=_to_float_safe(mapped.get("average_throughput_mbps")),
                storage_pool_name=mapped.get("storage_pool_name") or None,
                array_name=mapped.get("array_name") or None,
                array_vendor=mapped.get("array_vendor") or None,
                average_latency_ms=_to_float_safe(mapped.get("average_latency_ms")),
                application=mapped.get("application") or None,
            ))
            added += 1
        except ValueError as exc:
            errors.append(f"FileNAS row {i + 2}: {exc} - skipped.")
    return added, errors


def _import_block_storage_rows(rows_of_dicts, customer_id, gathering_request_id=None):
    added = 0
    errors = []
    for i, row in enumerate(rows_of_dicts):
        norm = _normalize_row(row)
        mapped = {}
        for raw_col, val in norm.items():
            field = BLOCK_STORAGE_FIELD_MAP.get(raw_col)
            if field:
                mapped[field] = val

        volume_name = mapped.get("volume_name", "")
        if not volume_name:
            continue
        try:
            db.session.add(GatheringBlockStorageDetail(
                customer_id=customer_id,
                gathering_request_id=gathering_request_id,
                volume_name=volume_name,
                total_used_capacity_gb=_to_float_safe(mapped.get("total_used_capacity_gb")),
                total_provisioned_capacity_gb=_to_float_safe(mapped.get("total_provisioned_capacity_gb")),
                peak_iops=_to_float_safe(mapped.get("peak_iops")),
                peak_throughput_mbps=_to_float_safe(mapped.get("peak_throughput_mbps")),
                average_iops=_to_float_safe(mapped.get("average_iops")),
                average_throughput_mbps=_to_float_safe(mapped.get("average_throughput_mbps")),
                array_name=mapped.get("array_name") or None,
                average_latency_ms=_to_float_safe(mapped.get("average_latency_ms")),
                application=mapped.get("application") or None,
            ))
            added += 1
        except ValueError as exc:
            errors.append(f"BlockStorage row {i + 2}: {exc} - skipped.")
    return added, errors


def _read_csv_rows(file_stream):
    content = file_stream.read().decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(content))
    return list(reader)


def _read_xlsx_rows_by_sheet(file_stream):
    import openpyxl  # noqa: PLC0415

    wb = openpyxl.load_workbook(file_stream, data_only=True)
    sheet_rows = {}
    for ws in wb.worksheets:
        row_iter = ws.iter_rows(min_row=1, max_row=1)
        first = next(row_iter, None)
        if not first:
            continue
        headers = [str(c.value or "").strip() for c in first]
        if not any(headers):
            continue

        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(v not in (None, "") for v in row):
                continue
            rows.append({headers[j]: ("" if row[j] is None else str(row[j])) for j in range(len(headers))})
        sheet_rows[ws.title.strip().lower()] = rows
    return sheet_rows


def _detect_rows_type(rows):
    if not rows:
        return "unknown"
    normalized_keys = {(str(k or "")).strip().lower() for k in rows[0].keys()}
    if "server name" in normalized_keys:
        return "server"
    if "file server/share name" in normalized_keys:
        return "file_nas"
    if "volume name" in normalized_keys:
        return "block_storage"
    return "unknown"


def _import_uploaded_file_to_all_tables(file_path, customer_id, gathering_request_id=None):
    """Import uploaded customer file (.csv/.xlsx) into relevant gathering tables."""
    p = Path(file_path)
    name = p.name.lower()

    total_server = 0
    total_file_nas = 0
    total_block = 0
    errors = []

    if name.endswith(".csv"):
        with p.open("rb") as f:
            rows = _read_csv_rows(f)
        detected_type = _detect_rows_type(rows)
        if detected_type == "server":
            total_server, errors = _import_server_rows(rows, customer_id, gathering_request_id)
        elif detected_type == "file_nas":
            total_file_nas, errors = _import_file_nas_rows(rows, customer_id, gathering_request_id)
        elif detected_type == "block_storage":
            total_block, errors = _import_block_storage_rows(rows, customer_id, gathering_request_id)
        else:
            errors.append("Could not detect CSV type.")
        return total_server, total_file_nas, total_block, errors

    if name.endswith(".xlsx") or name.endswith(".xls"):
        with p.open("rb") as f:
            sheets = _read_xlsx_rows_by_sheet(f)
        server_rows = sheets.get("template", [])
        file_nas_rows = sheets.get("filenas storage (if applicable)", [])
        block_rows = sheets.get("block storage (if applicable)", [])

        if not server_rows and not file_nas_rows and not block_rows:
            for _, rows in sheets.items():
                kind = _detect_rows_type(rows)
                if kind == "server":
                    server_rows.extend(rows)
                elif kind == "file_nas":
                    file_nas_rows.extend(rows)
                elif kind == "block_storage":
                    block_rows.extend(rows)

        s_added, s_err = _import_server_rows(server_rows, customer_id, gathering_request_id)
        f_added, f_err = _import_file_nas_rows(file_nas_rows, customer_id, gathering_request_id)
        b_added, b_err = _import_block_storage_rows(block_rows, customer_id, gathering_request_id)
        total_server += s_added
        total_file_nas += f_added
        total_block += b_added
        errors.extend(s_err + f_err + b_err)

    return total_server, total_file_nas, total_block, errors


def _import_uploaded_file_for_target(upload, target, customer_id, gathering_request_id=None):
    filename = (upload.filename or "").lower()
    rows = []
    normalized_target = (target or "").strip().lower()

    if filename.endswith(".csv"):
        rows = _read_csv_rows(upload.stream)
    elif filename.endswith(".xlsx") or filename.endswith(".xls"):
        sheets = _read_xlsx_rows_by_sheet(upload.stream)
        if normalized_target == "server":
            rows = sheets.get("template", [])
        elif normalized_target == "file-nas":
            rows = sheets.get("filenas storage (if applicable)", [])
        elif normalized_target == "block-storage":
            rows = sheets.get("block storage (if applicable)", [])

        if not rows and sheets:
            rows = next(iter(sheets.values()))
    else:
        return 0, ["Unsupported file format. Use .csv or .xlsx."]

    if normalized_target == "server":
        return _import_server_rows(rows, customer_id, gathering_request_id)
    if normalized_target == "file-nas":
        return _import_file_nas_rows(rows, customer_id, gathering_request_id)
    if normalized_target == "block-storage":
        return _import_block_storage_rows(rows, customer_id, gathering_request_id)
    return 0, ["Invalid import target."]
