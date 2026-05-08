import csv
import io

from flask import flash, redirect, request, send_file, url_for
from flask_login import login_required

from app.models import (
    Customer,
    GatheringBlockStorageDetail,
    GatheringFileNasDetail,
    GatheringServerDetail,
    db,
)


def register_gathering_data_routes(
    bp,
    *,
    detect_rows_type,
    import_block_storage_rows,
    import_file_nas_rows,
    import_server_rows,
    log_opportunity_history,
    read_csv_rows,
    read_xlsx_rows_by_sheet,
    to_float,
    to_int,
):
    if getattr(bp, "_gathering_data_routes_registered", False):
        return
    setattr(bp, "_gathering_data_routes_registered", True)

    @bp.route("/opportunities/<int:customer_id>/gathering-data/<int:entry_id>/edit", methods=["POST"])
    @login_required
    def opportunity_gathering_data_edit(customer_id, entry_id):
        customer = Customer.query.get_or_404(customer_id)
        entry = GatheringServerDetail.query.get_or_404(entry_id)
        if entry.customer_id != customer.id:
            flash("Invalid gathering data entry.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        try:
            server_name = (request.form.get("server_name") or "").strip()
            if not server_name:
                flash("Server name is required.", "error")
                return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

            entry.server_name = server_name
            entry.cpu_cores = to_int(request.form.get("cpu_cores"))
            entry.memory_mb = to_int(request.form.get("memory_mb"))
            entry.provisioned_storage_gb = to_float(request.form.get("provisioned_storage_gb"))
            entry.operating_system = (request.form.get("operating_system") or "").strip() or None
            entry.is_virtual = (request.form.get("is_virtual") or "").strip().lower() in ("yes", "true", "1")
            entry.hypervisor_name = (request.form.get("hypervisor_name") or "").strip() or None
            entry.cpu_string = (request.form.get("cpu_string") or "").strip() or None
            entry.environment = (request.form.get("environment") or "").strip() or None
            entry.sql_edition = (request.form.get("sql_edition") or "").strip() or None
            entry.application = (request.form.get("application") or "").strip() or None
            entry.storage_type = (request.form.get("storage_type") or "").strip() or None
            entry.cpu_utilization_peak = to_float(request.form.get("cpu_utilization_peak"))
            entry.memory_utilization_peak = to_float(request.form.get("memory_utilization_peak"))
            entry.time_in_use = to_float(request.form.get("time_in_use"))
            entry.annual_cost_usd = to_float(request.form.get("annual_cost_usd"))
            db.session.commit()

            log_opportunity_history(
                customer.id,
                action="gathering-data-edited",
                changes_summary=f"Gathering data entry updated for server '{entry.server_name}'.",
                remark=(request.form.get("change_remark") or "").strip() or None,
            )
            db.session.commit()
            flash("Gathering data entry updated.", "success")
        except ValueError:
            flash("Please enter valid numeric values for numeric fields.", "error")

        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/<int:customer_id>/gathering-data/<int:entry_id>/delete-from-edit", methods=["POST"])
    @login_required
    def opportunity_gathering_data_delete_from_edit(customer_id, entry_id):
        customer = Customer.query.get_or_404(customer_id)
        entry = GatheringServerDetail.query.get_or_404(entry_id)
        if entry.customer_id != customer.id:
            flash("Invalid gathering data entry.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        sname = entry.server_name
        db.session.delete(entry)
        log_opportunity_history(
            customer.id,
            action="gathering-data-deleted",
            changes_summary=f"Gathering data entry deleted for server '{sname}'.",
            remark="Deleted from opportunity edit page",
        )
        db.session.commit()
        flash("Gathering data entry deleted.", "success")
        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/<int:customer_id>/gathering-data/add", methods=["POST"])
    @login_required
    def opportunity_gathering_data_add(customer_id):
        customer = Customer.query.get_or_404(customer_id)
        try:
            server_name = (request.form.get("server_name") or "").strip()
            if not server_name:
                flash("Server name is required.", "error")
                return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

            entry = GatheringServerDetail(
                customer_id=customer.id,
                server_name=server_name,
                cpu_cores=to_int(request.form.get("cpu_cores")),
                memory_mb=to_int(request.form.get("memory_mb")),
                provisioned_storage_gb=to_float(request.form.get("provisioned_storage_gb")),
                operating_system=(request.form.get("operating_system") or "").strip() or None,
                is_virtual=(request.form.get("is_virtual") or "").strip().lower() in ("yes", "true", "1"),
                hypervisor_name=(request.form.get("hypervisor_name") or "").strip() or None,
                cpu_string=(request.form.get("cpu_string") or "").strip() or None,
                environment=(request.form.get("environment") or "").strip() or None,
                sql_edition=(request.form.get("sql_edition") or "").strip() or None,
                application=(request.form.get("application") or "").strip() or None,
                storage_type=(request.form.get("storage_type") or "").strip() or None,
                cpu_utilization_peak=to_float(request.form.get("cpu_utilization_peak")),
                memory_utilization_peak=to_float(request.form.get("memory_utilization_peak")),
                time_in_use=to_float(request.form.get("time_in_use")),
                annual_cost_usd=to_float(request.form.get("annual_cost_usd")),
            )
            db.session.add(entry)
            log_opportunity_history(
                customer.id,
                action="gathering-data-added",
                changes_summary=f"Gathering data entry added for server '{server_name}'.",
                remark=(request.form.get("change_remark") or "").strip() or None,
            )
            db.session.commit()
            flash("Gathering data row saved.", "success")
        except ValueError:
            flash("Please enter valid numeric values.", "error")
        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/<int:customer_id>/gathering-file-nas/add", methods=["POST"])
    @login_required
    def opportunity_gathering_file_nas_add(customer_id):
        customer = Customer.query.get_or_404(customer_id)
        try:
            name = (request.form.get("file_server_share_name") or "").strip()
            if not name:
                flash("File Server/Share Name is required.", "error")
                return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

            entry = GatheringFileNasDetail(
                customer_id=customer.id,
                file_server_share_name=name,
                total_used_capacity_gb=to_float(request.form.get("total_used_capacity_gb")),
                access_protocol=(request.form.get("access_protocol") or "").strip() or None,
                total_provisioned_capacity_gb=to_float(request.form.get("total_provisioned_capacity_gb")),
                storage_efficiency_ratio=to_float(request.form.get("storage_efficiency_ratio")),
                peak_iops=to_float(request.form.get("peak_iops")),
                peak_throughput_mbps=to_float(request.form.get("peak_throughput_mbps")),
                average_iops=to_float(request.form.get("average_iops")),
                average_throughput_mbps=to_float(request.form.get("average_throughput_mbps")),
                storage_pool_name=(request.form.get("storage_pool_name") or "").strip() or None,
                array_name=(request.form.get("array_name") or "").strip() or None,
                array_vendor=(request.form.get("array_vendor") or "").strip() or None,
                average_latency_ms=to_float(request.form.get("average_latency_ms")),
                application=(request.form.get("application") or "").strip() or None,
            )
            db.session.add(entry)
            log_opportunity_history(
                customer.id,
                action="gathering-file-nas-added",
                changes_summary=f"FileNAS entry added: '{name}'.",
                remark=(request.form.get("change_remark") or "").strip() or None,
            )
            db.session.commit()
            flash("FileNAS row saved.", "success")
        except ValueError:
            flash("Please enter valid numeric values.", "error")
        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/<int:customer_id>/gathering-file-nas/<int:entry_id>/edit", methods=["POST"])
    @login_required
    def opportunity_gathering_file_nas_edit(customer_id, entry_id):
        customer = Customer.query.get_or_404(customer_id)
        entry = GatheringFileNasDetail.query.get_or_404(entry_id)
        if entry.customer_id != customer.id:
            flash("Invalid FileNAS entry.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        try:
            name = (request.form.get("file_server_share_name") or "").strip()
            if not name:
                flash("File Server/Share Name is required.", "error")
                return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

            entry.file_server_share_name = name
            entry.total_used_capacity_gb = to_float(request.form.get("total_used_capacity_gb"))
            entry.access_protocol = (request.form.get("access_protocol") or "").strip() or None
            entry.total_provisioned_capacity_gb = to_float(request.form.get("total_provisioned_capacity_gb"))
            entry.storage_efficiency_ratio = to_float(request.form.get("storage_efficiency_ratio"))
            entry.peak_iops = to_float(request.form.get("peak_iops"))
            entry.peak_throughput_mbps = to_float(request.form.get("peak_throughput_mbps"))
            entry.average_iops = to_float(request.form.get("average_iops"))
            entry.average_throughput_mbps = to_float(request.form.get("average_throughput_mbps"))
            entry.storage_pool_name = (request.form.get("storage_pool_name") or "").strip() or None
            entry.array_name = (request.form.get("array_name") or "").strip() or None
            entry.array_vendor = (request.form.get("array_vendor") or "").strip() or None
            entry.average_latency_ms = to_float(request.form.get("average_latency_ms"))
            entry.application = (request.form.get("application") or "").strip() or None
            db.session.commit()

            log_opportunity_history(
                customer.id,
                action="gathering-file-nas-edited",
                changes_summary=f"FileNAS entry updated: '{name}'.",
                remark=(request.form.get("change_remark") or "").strip() or None,
            )
            db.session.commit()
            flash("FileNAS row updated.", "success")
        except ValueError:
            flash("Please enter valid numeric values.", "error")

        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/<int:customer_id>/gathering-file-nas/<int:entry_id>/delete", methods=["POST"])
    @login_required
    def opportunity_gathering_file_nas_delete(customer_id, entry_id):
        customer = Customer.query.get_or_404(customer_id)
        entry = GatheringFileNasDetail.query.get_or_404(entry_id)
        if entry.customer_id != customer.id:
            flash("Invalid FileNAS entry.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        name = entry.file_server_share_name
        db.session.delete(entry)
        log_opportunity_history(
            customer.id,
            action="gathering-file-nas-deleted",
            changes_summary=f"FileNAS entry deleted: '{name}'.",
            remark="Deleted from opportunity edit page",
        )
        db.session.commit()
        flash("FileNAS row deleted.", "success")
        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/<int:customer_id>/gathering-block-storage/add", methods=["POST"])
    @login_required
    def opportunity_gathering_block_storage_add(customer_id):
        customer = Customer.query.get_or_404(customer_id)
        try:
            name = (request.form.get("volume_name") or "").strip()
            if not name:
                flash("Volume Name is required.", "error")
                return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

            entry = GatheringBlockStorageDetail(
                customer_id=customer.id,
                volume_name=name,
                total_used_capacity_gb=to_float(request.form.get("total_used_capacity_gb")),
                total_provisioned_capacity_gb=to_float(request.form.get("total_provisioned_capacity_gb")),
                peak_iops=to_float(request.form.get("peak_iops")),
                peak_throughput_mbps=to_float(request.form.get("peak_throughput_mbps")),
                average_iops=to_float(request.form.get("average_iops")),
                average_throughput_mbps=to_float(request.form.get("average_throughput_mbps")),
                array_name=(request.form.get("array_name") or "").strip() or None,
                average_latency_ms=to_float(request.form.get("average_latency_ms")),
                application=(request.form.get("application") or "").strip() or None,
            )
            db.session.add(entry)
            log_opportunity_history(
                customer.id,
                action="gathering-block-storage-added",
                changes_summary=f"Block Storage entry added: '{name}'.",
                remark=(request.form.get("change_remark") or "").strip() or None,
            )
            db.session.commit()
            flash("Block Storage row saved.", "success")
        except ValueError:
            flash("Please enter valid numeric values.", "error")
        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/<int:customer_id>/gathering-block-storage/<int:entry_id>/edit", methods=["POST"])
    @login_required
    def opportunity_gathering_block_storage_edit(customer_id, entry_id):
        customer = Customer.query.get_or_404(customer_id)
        entry = GatheringBlockStorageDetail.query.get_or_404(entry_id)
        if entry.customer_id != customer.id:
            flash("Invalid Block Storage entry.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        try:
            name = (request.form.get("volume_name") or "").strip()
            if not name:
                flash("Volume Name is required.", "error")
                return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

            entry.volume_name = name
            entry.total_used_capacity_gb = to_float(request.form.get("total_used_capacity_gb"))
            entry.total_provisioned_capacity_gb = to_float(request.form.get("total_provisioned_capacity_gb"))
            entry.peak_iops = to_float(request.form.get("peak_iops"))
            entry.peak_throughput_mbps = to_float(request.form.get("peak_throughput_mbps"))
            entry.average_iops = to_float(request.form.get("average_iops"))
            entry.average_throughput_mbps = to_float(request.form.get("average_throughput_mbps"))
            entry.array_name = (request.form.get("array_name") or "").strip() or None
            entry.average_latency_ms = to_float(request.form.get("average_latency_ms"))
            entry.application = (request.form.get("application") or "").strip() or None
            db.session.commit()

            log_opportunity_history(
                customer.id,
                action="gathering-block-storage-edited",
                changes_summary=f"Block Storage entry updated: '{name}'.",
                remark=(request.form.get("change_remark") or "").strip() or None,
            )
            db.session.commit()
            flash("Block Storage row updated.", "success")
        except ValueError:
            flash("Please enter valid numeric values.", "error")

        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/<int:customer_id>/gathering-block-storage/<int:entry_id>/delete", methods=["POST"])
    @login_required
    def opportunity_gathering_block_storage_delete(customer_id, entry_id):
        customer = Customer.query.get_or_404(customer_id)
        entry = GatheringBlockStorageDetail.query.get_or_404(entry_id)
        if entry.customer_id != customer.id:
            flash("Invalid Block Storage entry.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        name = entry.volume_name
        db.session.delete(entry)
        log_opportunity_history(
            customer.id,
            action="gathering-block-storage-deleted",
            changes_summary=f"Block Storage entry deleted: '{name}'.",
            remark="Deleted from opportunity edit page",
        )
        db.session.commit()
        flash("Block Storage row deleted.", "success")
        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/<int:customer_id>/gathering-data/import/<string:import_target>", methods=["POST"])
    @login_required
    def opportunity_gathering_data_import_single(customer_id, import_target):
        customer = Customer.query.get_or_404(customer_id)
        upload = request.files.get("import_file")
        if not upload or not upload.filename:
            flash("Please select a CSV or XLSX file to import.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        target = (import_target or "").strip().lower()
        if target not in ("server", "file-nas", "block-storage"):
            flash("Invalid import target.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        filename = upload.filename.lower()
        rows = []
        if filename.endswith(".csv"):
            rows = read_csv_rows(upload.stream)
        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            try:
                sheets = read_xlsx_rows_by_sheet(upload.stream)
                if target == "server":
                    rows = sheets.get("template", [])
                elif target == "file-nas":
                    rows = sheets.get("filenas storage (if applicable)", [])
                elif target == "block-storage":
                    rows = sheets.get("block storage (if applicable)", [])

                if not rows and sheets:
                    rows = next(iter(sheets.values()))
            except ImportError:
                flash("openpyxl is not installed. Run: pip install openpyxl", "error")
                return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))
        else:
            flash("Unsupported file format. Use .csv or .xlsx.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        added = 0
        errors = []
        if target == "server":
            added, errors = import_server_rows(rows, customer.id)
        elif target == "file-nas":
            added, errors = import_file_nas_rows(rows, customer.id)
        elif target == "block-storage":
            added, errors = import_block_storage_rows(rows, customer.id)

        if added:
            log_opportunity_history(
                customer.id,
                action="gathering-data-imported",
                changes_summary=f"Imported {added} row(s) for {target} from single-sheet file.",
            )
            db.session.commit()
            flash(f"Imported {added} row(s) for {target}.", "success")
        else:
            flash(f"No valid rows were imported for {target}.", "error")

        if errors:
            for e in errors[:6]:
                flash(e, "error")

        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/opportunities/gathering-data/sample.csv")
    @login_required
    def gathering_data_sample_csv():
        columns = [
            "Server Name", "CPU Cores", "Memory (MB)", "Provisioned Storage (GB)",
            "Operating System", "Is Virtual (Yes/No)", "Hypervisor Name", "CPU String",
            "Environment", "SQL Edition", "Application",
            "CPU Util Peak (%)", "Memory Util Peak (%)", "Time In-Use (%)",
            "Annual Cost (USD)", "Storage Type",
        ]
        sample = [
            "Apache01", "4", "4096", "500",
            "Windows Server 2012 R2", "Yes", "Host-1", "Intel Xeon E7-8893 v4 @ 3.2GHz",
            "Production", "SQL Server 2012 Enterprise", "Service Now",
            "0.60", "0.95", "1.00",
            "3400", "HDD",
        ]
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(columns)
        writer.writerow(sample)
        output.seek(0)
        return send_file(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            mimetype="text/csv",
            as_attachment=True,
            download_name="gathering_data_sample.csv",
        )

    @bp.route("/opportunities/gathering-data/sample-file-nas.csv")
    @login_required
    def gathering_data_sample_file_nas_csv():
        columns = [
            "File Server/Share Name",
            "Total Used Capacity (GB) - Usable",
            "Access Protocol (CIFS/NFS)",
            "Total Provisioned Capacity (GB) - Usable",
            "Storage Efficiency Ratio (Dedupe/Compression)",
            "Peak  IOPS (reads & writes)",
            "Peak Throughput (MBps)",
            "Average  IOPS (reads & writes)",
            "Average Throughput (MBps)",
            "Storage Pool Name",
            "Array Name",
            "Array Vendor",
            "Average Latency (ms)",
            "Application",
        ]
        sample = [
            "CorpFiles01", "820", "CIFS", "1200", "1.35", "14500", "1200", "9100", "760",
            "Pool-A", "NetApp01", "NetApp", "1.8", "File Sharing",
        ]
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(columns)
        writer.writerow(sample)
        output.seek(0)
        return send_file(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            mimetype="text/csv",
            as_attachment=True,
            download_name="gathering_file_nas_sample.csv",
        )

    @bp.route("/opportunities/gathering-data/sample-block-storage.csv")
    @login_required
    def gathering_data_sample_block_storage_csv():
        columns = [
            "Volume Name",
            "Total Used Capacity (GB) - Usable",
            "Total Provisioned Capacity (GB) - Usable",
            "Peak  IOPS (reads & writes)",
            "Peak Throughput (MBps)",
            "Average  IOPS (reads & writes)",
            "Average Throughput (MBps)",
            "Array Name",
            "Average Latency (ms)",
            "Application",
        ]
        sample = [
            "vol-oracle-prod", "2600", "3200", "42000", "1800", "30000", "1260", "PureArray01", "1.2", "Oracle DB",
        ]
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(columns)
        writer.writerow(sample)
        output.seek(0)
        return send_file(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            mimetype="text/csv",
            as_attachment=True,
            download_name="gathering_block_storage_sample.csv",
        )

    @bp.route("/opportunities/gathering-data/sample.xlsx")
    @login_required
    def gathering_data_sample_xlsx():
        try:
            import openpyxl  # noqa: PLC0415
        except ImportError:
            flash("openpyxl is not installed. Run: pip install openpyxl", "error")
            return redirect(url_for("crm.opportunity_list"))

        wb = openpyxl.Workbook()

        ws_instructions = wb.active
        ws_instructions.title = "Instructions"
        ws_instructions.append([
            "This workbook is the data import template for AWS migration assessment. Fill the applicable sheets and upload from Main Import."
        ])

        ws_glossary = wb.create_sheet("Glossary")
        ws_glossary.append(["Attribute Name", "Example", "Requirement", "Notes"])
        ws_glossary.append(["Server Name", "Apache01", "Required for Template sheet", "Unique workload identifier"])
        ws_glossary.append(["File Server/Share Name", "CorpFiles01", "Required for FileNAS rows", "Used for NAS import mapping"])
        ws_glossary.append(["Volume Name", "vol-oracle-prod", "Required for Block Storage rows", "Used for block import mapping"])

        ws_server = wb.create_sheet("Template")
        ws_server.append([
            "Server Name", "CPU Cores", "Memory (MB)", "Provisioned Storage (GB)",
            "Operating System", "Is Virtual?", "Hypervisor Name", "Cpu String",
            "Environment", "SQL Edition", "Application",
            "Cpu Utilization Peak (%)", "Memory Utilization Peak (%)", "Time In-Use (%)",
            "Annual Cost (USD)", "Storage Type",
        ])
        ws_server.append([
            "Apache01", 4, 4096, 500,
            "Windows Server 2012 R2", "Yes", "Host-1", "Intel Xeon E7-8893 v4 @ 3.2GHz",
            "Production", "SQL Server 2012 Enterprise", "Service Now",
            0.60, 0.95, 1.00,
            3400, "HDD",
        ])

        ws_file = wb.create_sheet("FileNAS Storage (If applicable)")
        ws_file.append([
            "File Server/Share Name", "Total Used Capacity (GB) - Usable", "Access Protocol (CIFS/NFS)",
            "Total Provisioned Capacity (GB) - Usable", "Storage Efficiency Ratio (Dedupe/Compression)",
            "Peak  IOPS (reads & writes)", "Peak Throughput (MBps)", "Average  IOPS (reads & writes)",
            "Average Throughput (MBps)", "Storage Pool Name", "Array Name", "Array Vendor",
            "Average Latency (ms)", "Application",
        ])
        ws_file.append([
            "CorpFiles01", 820, "CIFS", 1200, 1.35,
            14500, 1200, 9100, 760, "Pool-A", "NetApp01", "NetApp", 1.8, "File Sharing",
        ])

        ws_block = wb.create_sheet("Block Storage (If applicable)")
        ws_block.append([
            "Volume Name", "Total Used Capacity (GB) - Usable", "Total Provisioned Capacity (GB) - Usable",
            "Peak  IOPS (reads & writes)", "Peak Throughput (MBps)", "Average  IOPS (reads & writes)",
            "Average Throughput (MBps)", "Array Name", "Average Latency (ms)", "Application",
        ])
        ws_block.append([
            "vol-oracle-prod", 2600, 3200, 42000, 1800, 30000, 1260, "PureArray01", 1.2, "Oracle DB",
        ])

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="gathering_data_sample.xlsx",
        )

    @bp.route("/opportunities/<int:customer_id>/gathering-data/import", methods=["POST"])
    @login_required
    def opportunity_gathering_data_import(customer_id):
        customer = Customer.query.get_or_404(customer_id)
        upload = request.files.get("import_file")
        if not upload or not upload.filename:
            flash("Please select a CSV or XLSX file to import.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        filename = upload.filename.lower()
        total_server = 0
        total_file_nas = 0
        total_block = 0
        errors = []

        if filename.endswith(".csv"):
            rows = read_csv_rows(upload.stream)
            detected_type = detect_rows_type(rows)
            if detected_type == "server":
                total_server, errors = import_server_rows(rows, customer.id)
            elif detected_type == "file_nas":
                total_file_nas, errors = import_file_nas_rows(rows, customer.id)
            elif detected_type == "block_storage":
                total_block, errors = import_block_storage_rows(rows, customer.id)
            else:
                flash("Could not detect CSV type. Use Server, FileNAS, or Block Storage columns.", "error")
                return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            try:
                sheets = read_xlsx_rows_by_sheet(upload.stream)
            except ImportError:
                flash("openpyxl is not installed. Run: pip install openpyxl", "error")
                return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

            server_rows = sheets.get("template", [])
            file_nas_rows = sheets.get("filenas storage (if applicable)", [])
            block_rows = sheets.get("block storage (if applicable)", [])

            if not server_rows and not file_nas_rows and not block_rows:
                for _, rows in sheets.items():
                    kind = detect_rows_type(rows)
                    if kind == "server":
                        server_rows.extend(rows)
                    elif kind == "file_nas":
                        file_nas_rows.extend(rows)
                    elif kind == "block_storage":
                        block_rows.extend(rows)

            s_added, s_err = import_server_rows(server_rows, customer.id)
            f_added, f_err = import_file_nas_rows(file_nas_rows, customer.id)
            b_added, b_err = import_block_storage_rows(block_rows, customer.id)
            total_server += s_added
            total_file_nas += f_added
            total_block += b_added
            errors.extend(s_err + f_err + b_err)
        else:
            flash("Unsupported file format. Use .csv or .xlsx.", "error")
            return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

        total_added = total_server + total_file_nas + total_block
        if total_added:
            log_opportunity_history(
                customer.id,
                action="gathering-data-imported",
                changes_summary=(
                    f"Imported Server: {total_server}, FileNAS: {total_file_nas}, "
                    f"Block Storage: {total_block} row(s)."
                ),
            )
            db.session.commit()
            flash(
                f"Imported rows - Server: {total_server}, FileNAS: {total_file_nas}, Block Storage: {total_block}.",
                "success",
            )
        else:
            flash("No valid rows were imported.", "error")

        if errors:
            for e in errors[:8]:
                flash(e, "error")

        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))
