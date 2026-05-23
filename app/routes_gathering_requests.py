import csv
import io
import secrets
from datetime import datetime, timedelta
from pathlib import Path

from flask import (
    current_app,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from flask_login import login_required
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

from app.models import (
    Customer,
    EmailLog,
    GatheringBlockStorageDetail,
    GatheringFileNasDetail,
    GatheringRequest,
    GatheringServerDetail,
    db,
)
from app.services.email_service import EmailService
from app.services.macro_service import build_macro_values


def register_gathering_request_routes(
    bp,
    *,
    actor_name,
    generate_access_key,
    get_effective_app_base_url,
    import_uploaded_file_for_target,
    import_uploaded_file_to_all_tables,
    is_gathering_request_expired,
    log_opportunity_history,
    render_system_email_template,
):
    if getattr(bp, "_gathering_request_routes_registered", False):
        return
    setattr(bp, "_gathering_request_routes_registered", True)

    @bp.route("/gathering/send", methods=["GET", "POST"])
    def send_gathering_request():
        selected_customer_id = request.values.get("customer_id", type=int)
        customer_q = (request.values.get("customer_q") or "").strip()
        customer_page = request.values.get("customer_page", 1, type=int)
        customers_query = Customer.query
        if customer_q:
            like = f"%{customer_q}%"
            customers_query = customers_query.filter(
                db.or_(
                    Customer.customer_name.ilike(like),
                    Customer.email.ilike(like),
                    Customer.account_name.ilike(like),
                )
            )
        customers_pagination = customers_query.order_by(Customer.customer_name.asc()).paginate(
            page=max(customer_page or 1, 1), per_page=50, error_out=False
        )
        customers = list(customers_pagination.items)
        if selected_customer_id and all(c.id != selected_customer_id for c in customers):
            selected_customer = Customer.query.get(selected_customer_id)
            if selected_customer:
                customers.insert(0, selected_customer)
        default_expiry_days = 7
        form_values = {
            "cc": "",
            "bcc": "",
            "note": "",
            "expiry_days": default_expiry_days,
        }

        if request.method == "POST":
            customer_id = request.form.get("customer_id", type=int)
            note = (request.form.get("note") or "").strip()
            cc_raw = (request.form.get("cc_emails") or "").strip()
            bcc_raw = (request.form.get("bcc_emails") or "").strip()
            expiry_days = request.form.get("expiry_days", type=int) or default_expiry_days
            customer = Customer.query.get(customer_id)

            form_values = {
                "cc": cc_raw,
                "bcc": bcc_raw,
                "note": note,
                "expiry_days": expiry_days,
            }

            cc_list = [e.strip() for e in cc_raw.split(",") if e.strip()]
            bcc_list = [e.strip() for e in bcc_raw.split(",") if e.strip()]

            invalid_cc = [e for e in cc_list if "@" not in e]
            invalid_bcc = [e for e in bcc_list if "@" not in e]
            if invalid_cc or invalid_bcc:
                flash("Please enter valid CC/BCC emails separated by commas.", "error")
                return render_template(
                    "send_gathering.html",
                    customers=customers,
                    customers_pagination=customers_pagination,
                    customer_q=customer_q,
                    selected_customer_id=customer_id,
                    default_expiry_days=default_expiry_days,
                    form_values=form_values,
                )

            if not customer:
                flash("Select a valid customer.", "error")
                return render_template(
                    "send_gathering.html",
                    customers=customers,
                    customers_pagination=customers_pagination,
                    customer_q=customer_q,
                    selected_customer_id=selected_customer_id,
                    default_expiry_days=default_expiry_days,
                    form_values=form_values,
                )

            if expiry_days < 1 or expiry_days > 90:
                flash("Expiry must be between 1 and 90 days.", "error")
                return render_template(
                    "send_gathering.html",
                    customers=customers,
                    customers_pagination=customers_pagination,
                    customer_q=customer_q,
                    selected_customer_id=customer_id,
                    default_expiry_days=default_expiry_days,
                    form_values=form_values,
                )

            token = secrets.token_urlsafe(24)
            access_key = generate_access_key()
            attachment = request.files.get("gathering_sheet")
            saved_name = None

            if attachment and attachment.filename:
                safe_name = secure_filename(attachment.filename)
                saved_name = f"sheet_{datetime.utcnow().strftime('%Y%m%d%H%M%S')}_{safe_name}"
                target = Path(current_app.config["UPLOAD_FOLDER"]) / saved_name
                attachment.save(target)

            request_record = GatheringRequest(
                token=token,
                customer_id=customer.id,
                note=note,
                expires_at=datetime.utcnow() + timedelta(days=expiry_days),
                access_key_hash=generate_password_hash(access_key),
                access_key_hint=access_key[-4:],
                sheet_file_name=saved_name,
                status="sent",
            )
            db.session.add(request_record)
            db.session.flush()

            form_link = f"{get_effective_app_base_url()}/gathering/form/{token}"
            expires_at_text = request_record.expires_at.strftime("%Y-%m-%d %H:%M UTC") if request_record.expires_at else ""
            macro_values = build_macro_values(
                customer,
                {
                    "gathering_form_link": form_link,
                    "gathering_access_key": access_key,
                    "gathering_expires_at": expires_at_text,
                    "note": note,
                },
            )

            tpl, rendered_subject, rendered_body = render_system_email_template("gathering_request", macro_values)

            email_service = EmailService(current_app)
            result = email_service.send_html_email(
                customer.email,
                rendered_subject,
                rendered_body,
                cc_emails=cc_list,
                bcc_emails=bcc_list,
            )

            email_log = EmailLog(
                customer_id=customer.id,
                template_id=tpl.id if tpl else None,
                recipient_email=customer.email,
                email_type="gathering",
                subject=rendered_subject,
                body=rendered_body,
                status=result.status,
                error_message=result.error,
            )
            db.session.add(email_log)
            log_opportunity_history(
                customer.id,
                action="gathering-link-sent",
                changes_summary=(
                    f"Gathering request processed to {customer.email}. Subject: {rendered_subject}. "
                    f"Expiry: {expires_at_text}. CC: {', '.join(cc_list) if cc_list else 'None'}. "
                    f"BCC: {', '.join(bcc_list) if bcc_list else 'None'}. Status: {result.status}."
                ),
                remark=note or None,
                tag_name="gathering",
            )
            db.session.commit()

            if result.success:
                flash(f"Gathering request processed with status: {result.status}", "success")
            else:
                flash(f"Gathering email failed: {result.error}", "error")
            flash(
                f"Secure form link generated. Access key (save now): {access_key} | Expires: {expires_at_text}",
                "success",
            )

            return redirect(url_for("crm.send_gathering_request", customer_id=customer.id, customer_q=customer_q, customer_page=customer_page))

        return render_template(
            "send_gathering.html",
            customers=customers,
            customers_pagination=customers_pagination,
            customer_q=customer_q,
            selected_customer_id=selected_customer_id,
            default_expiry_days=default_expiry_days,
            form_values=form_values,
        )

    @bp.route("/gathering/form/<token>", methods=["GET", "POST"])
    def gathering_form(token):
        request_record = GatheringRequest.query.filter_by(token=token).first_or_404()
        is_expired = is_gathering_request_expired(request_record)
        session_key = f"gathering_access_ok_{request_record.id}"
        key_required = bool(request_record.access_key_hash)
        key_verified = bool(session.get(session_key))

        # Locked by admin - block all edits and imports.
        if request_record.is_locked:
            return render_template(
                "gathering_form.html",
                request_record=request_record,
                submitted=False,
                is_locked=True,
            )

        if is_expired:
            if request_record.status not in ("submitted", "expired"):
                request_record.status = "expired"
                db.session.commit()
            return render_template(
                "gathering_form.html",
                request_record=request_record,
                submitted=False,
                link_expired=True,
                key_required=False,
            )

        if request.method == "POST" and request.form.get("form_action") == "verify_key":
            entered_key = (request.form.get("access_key") or "").strip().upper()
            if not entered_key:
                return render_template(
                    "gathering_form.html",
                    request_record=request_record,
                    submitted=False,
                    key_required=True,
                    key_error="Enter the access key shared in the email.",
                )

            if not request_record.access_key_hash or check_password_hash(request_record.access_key_hash, entered_key):
                session[session_key] = True
                request_record.access_verified_at = request_record.access_verified_at or datetime.utcnow()
                db.session.commit()
                return redirect(url_for("crm.gathering_form", token=token))

            return render_template(
                "gathering_form.html",
                request_record=request_record,
                submitted=False,
                key_required=True,
                key_error="Invalid key. Please check the email and try again.",
            )

        if key_required and not key_verified:
            return render_template(
                "gathering_form.html",
                request_record=request_record,
                submitted=False,
                key_required=True,
            )

        if request.method == "POST" and request.form.get("form_action") in (
            "import-main",
            "import-server",
            "import-file-nas",
            "import-block-storage",
        ):
            action = request.form.get("form_action")
            import_errors = []
            import_message = None

            if action == "import-main":
                upload = request.files.get("submitted_sheet")
                if not upload or not upload.filename:
                    import_errors.append("Please select a workbook/CSV file for main import.")
                else:
                    safe_name = secure_filename(upload.filename)
                    final_name = f"submitted_import_{request_record.id}_{safe_name}"
                    target = Path(current_app.config["UPLOAD_FOLDER"]) / final_name
                    upload.save(target)
                    try:
                        s_added, f_added, b_added, parse_errors = import_uploaded_file_to_all_tables(
                            str(target),
                            request_record.customer_id,
                            request_record.id,
                        )
                        import_errors.extend(parse_errors)
                        imported_total = s_added + f_added + b_added
                        if imported_total:
                            log_opportunity_history(
                                request_record.customer_id,
                                action="gathering-data-imported",
                                changes_summary=(
                                    f"Customer imported workbook - Server: {s_added}, "
                                    f"FileNAS: {f_added}, Block Storage: {b_added}."
                                ),
                            )
                            db.session.commit()
                            import_message = (
                                f"Imported rows - Server: {s_added}, FileNAS: {f_added}, "
                                f"Block Storage: {b_added}."
                            )
                        else:
                            import_errors.append("No valid rows found in the selected file.")
                    except Exception:
                        import_errors.append("Could not parse file. Please upload a valid CSV/XLSX.")
                    finally:
                        try:
                            target.unlink(missing_ok=True)
                        except OSError:
                            pass
            else:
                upload_map = {
                    "import-server": ("submitted_server_sheet", "server", "Server"),
                    "import-file-nas": ("submitted_file_nas_sheet", "file-nas", "FileNAS"),
                    "import-block-storage": ("submitted_block_storage_sheet", "block-storage", "Block Storage"),
                }
                field_name, target_type, target_label = upload_map[action]
                upload = request.files.get(field_name)
                if not upload or not upload.filename:
                    import_errors.append(f"Please select a file for {target_label} import.")
                else:
                    try:
                        added, parse_errors = import_uploaded_file_for_target(
                            upload,
                            target_type,
                            request_record.customer_id,
                            request_record.id,
                        )
                        import_errors.extend(parse_errors)
                        if added:
                            log_opportunity_history(
                                request_record.customer_id,
                                action="gathering-data-imported",
                                changes_summary=f"Customer imported {added} {target_label} row(s).",
                            )
                            db.session.commit()
                            import_message = f"Imported {added} {target_label} row(s)."
                        else:
                            import_errors.append(f"No valid {target_label} rows found in file.")
                    except Exception:
                        import_errors.append(f"Could not parse {target_label} import file.")

            return render_template(
                "gathering_form.html",
                request_record=request_record,
                submitted=False,
                import_message=import_message,
                import_errors=import_errors,
            )

        if request.method == "POST":
            request_record.company_website = (request.form.get("company_website") or "").strip()
            request_record.current_tools = (request.form.get("current_tools") or "").strip()
            request_record.pain_points = (request.form.get("pain_points") or "").strip()

            upload = request.files.get("submitted_sheet")
            if upload and upload.filename:
                safe_name = secure_filename(upload.filename)
                final_name = f"submitted_{request_record.id}_{safe_name}"
                target = Path(current_app.config["UPLOAD_FOLDER"]) / final_name
                upload.save(target)
                request_record.submitted_sheet_path = str(target)

                # If customer uploads CSV/XLSX gathering sheet, auto-import all records.
                try:
                    s_added, f_added, b_added, _import_errors = import_uploaded_file_to_all_tables(
                        str(target),
                        request_record.customer_id,
                        request_record.id,
                    )
                    imported_total = s_added + f_added + b_added
                    if imported_total:
                        log_opportunity_history(
                            request_record.customer_id,
                            action="gathering-data-imported",
                            changes_summary=(
                                f"Customer uploaded sheet import - Server: {s_added}, "
                                f"FileNAS: {f_added}, Block Storage: {b_added}."
                            ),
                        )
                except Exception:
                    # Keep form submission resilient even if import parsing fails.
                    pass

            # Save structured server rows submitted by the customer.
            server_names = request.form.getlist("server_name[]")
            cpu_cores_list = request.form.getlist("cpu_cores[]")
            memory_mb_list = request.form.getlist("memory_mb[]")
            storage_list = request.form.getlist("provisioned_storage_gb[]")
            os_list = request.form.getlist("operating_system[]")
            virtual_list = request.form.getlist("is_virtual[]")
            hypervisor_list = request.form.getlist("hypervisor_name[]")
            cpu_str_list = request.form.getlist("cpu_string[]")
            env_list = request.form.getlist("environment[]")
            sql_list = request.form.getlist("sql_edition[]")
            app_list = request.form.getlist("application[]")
            stype_list = request.form.getlist("storage_type[]")
            cpu_util_list = request.form.getlist("cpu_utilization_peak[]")
            mem_util_list = request.form.getlist("memory_utilization_peak[]")
            tiu_list = request.form.getlist("time_in_use[]")
            cost_list = request.form.getlist("annual_cost_usd[]")

            def _gf(lst, i):
                try:
                    v = lst[i].strip() if i < len(lst) else ""
                    return float(v) if v else None
                except (ValueError, IndexError):
                    return None

            def _gi(lst, i):
                try:
                    v = lst[i].strip() if i < len(lst) else ""
                    return int(v) if v else None
                except (ValueError, IndexError):
                    return None

            def _gs(lst, i):
                return (lst[i] or "").strip() or None if i < len(lst) else None

            for idx, sname in enumerate(server_names):
                if not (sname or "").strip():
                    continue
                db.session.add(GatheringServerDetail(
                    customer_id=request_record.customer_id,
                    gathering_request_id=request_record.id,
                    server_name=sname.strip(),
                    cpu_cores=_gi(cpu_cores_list, idx),
                    memory_mb=_gi(memory_mb_list, idx),
                    provisioned_storage_gb=_gf(storage_list, idx),
                    operating_system=_gs(os_list, idx),
                    is_virtual=(_gs(virtual_list, idx) or "").lower() in ("yes", "true", "1"),
                    hypervisor_name=_gs(hypervisor_list, idx),
                    cpu_string=_gs(cpu_str_list, idx),
                    environment=_gs(env_list, idx),
                    sql_edition=_gs(sql_list, idx),
                    application=_gs(app_list, idx),
                    storage_type=_gs(stype_list, idx),
                    cpu_utilization_peak=_gf(cpu_util_list, idx),
                    memory_utilization_peak=_gf(mem_util_list, idx),
                    time_in_use=_gf(tiu_list, idx),
                    annual_cost_usd=_gf(cost_list, idx),
                ))

            # Save FileNAS rows.
            file_share_names = request.form.getlist("file_server_share_name[]")
            file_total_used = request.form.getlist("file_total_used_capacity_gb[]")
            file_access_protocol = request.form.getlist("file_access_protocol[]")
            file_total_provisioned = request.form.getlist("file_total_provisioned_capacity_gb[]")
            file_storage_eff_ratio = request.form.getlist("file_storage_efficiency_ratio[]")
            file_peak_iops = request.form.getlist("file_peak_iops[]")
            file_peak_throughput = request.form.getlist("file_peak_throughput_mbps[]")
            file_avg_iops = request.form.getlist("file_average_iops[]")
            file_avg_throughput = request.form.getlist("file_average_throughput_mbps[]")
            file_storage_pool_name = request.form.getlist("file_storage_pool_name[]")
            file_array_name = request.form.getlist("file_array_name[]")
            file_array_vendor = request.form.getlist("file_array_vendor[]")
            file_avg_latency = request.form.getlist("file_average_latency_ms[]")
            file_application = request.form.getlist("file_application[]")

            for idx, name in enumerate(file_share_names):
                if not (name or "").strip():
                    continue
                db.session.add(GatheringFileNasDetail(
                    customer_id=request_record.customer_id,
                    gathering_request_id=request_record.id,
                    file_server_share_name=name.strip(),
                    total_used_capacity_gb=_gf(file_total_used, idx),
                    access_protocol=_gs(file_access_protocol, idx),
                    total_provisioned_capacity_gb=_gf(file_total_provisioned, idx),
                    storage_efficiency_ratio=_gf(file_storage_eff_ratio, idx),
                    peak_iops=_gf(file_peak_iops, idx),
                    peak_throughput_mbps=_gf(file_peak_throughput, idx),
                    average_iops=_gf(file_avg_iops, idx),
                    average_throughput_mbps=_gf(file_avg_throughput, idx),
                    storage_pool_name=_gs(file_storage_pool_name, idx),
                    array_name=_gs(file_array_name, idx),
                    array_vendor=_gs(file_array_vendor, idx),
                    average_latency_ms=_gf(file_avg_latency, idx),
                    application=_gs(file_application, idx),
                ))

            # Save Block Storage rows.
            block_volume_name = request.form.getlist("block_volume_name[]")
            block_total_used = request.form.getlist("block_total_used_capacity_gb[]")
            block_total_provisioned = request.form.getlist("block_total_provisioned_capacity_gb[]")
            block_peak_iops = request.form.getlist("block_peak_iops[]")
            block_peak_throughput = request.form.getlist("block_peak_throughput_mbps[]")
            block_avg_iops = request.form.getlist("block_average_iops[]")
            block_avg_throughput = request.form.getlist("block_average_throughput_mbps[]")
            block_array_name = request.form.getlist("block_array_name[]")
            block_avg_latency = request.form.getlist("block_average_latency_ms[]")
            block_application = request.form.getlist("block_application[]")

            for idx, name in enumerate(block_volume_name):
                if not (name or "").strip():
                    continue
                db.session.add(GatheringBlockStorageDetail(
                    customer_id=request_record.customer_id,
                    gathering_request_id=request_record.id,
                    volume_name=name.strip(),
                    total_used_capacity_gb=_gf(block_total_used, idx),
                    total_provisioned_capacity_gb=_gf(block_total_provisioned, idx),
                    peak_iops=_gf(block_peak_iops, idx),
                    peak_throughput_mbps=_gf(block_peak_throughput, idx),
                    average_iops=_gf(block_avg_iops, idx),
                    average_throughput_mbps=_gf(block_avg_throughput, idx),
                    array_name=_gs(block_array_name, idx),
                    average_latency_ms=_gf(block_avg_latency, idx),
                    application=_gs(block_application, idx),
                ))

            if request_record.status not in ("submitted", "locked"):
                request_record.status = "submitted"
            request_record.submitted_at = datetime.utcnow()
            db.session.commit()

            return redirect(url_for("crm.gathering_form", token=token, saved=1))

        saved = request.args.get("saved")
        return render_template(
            "gathering_form.html",
            request_record=request_record,
            submitted=False,
            just_saved=bool(saved),
        )

    @bp.route("/gathering/request/<int:request_id>/lock", methods=["POST"])
    @login_required
    def gathering_request_lock(request_id):
        gr = GatheringRequest.query.get_or_404(request_id)
        gr.is_locked = True
        gr.locked_at = datetime.utcnow()
        gr.locked_by = actor_name()
        log_opportunity_history(
            gr.customer_id,
            action="gathering-locked",
            changes_summary=f"Gathering request #{gr.id} locked by {gr.locked_by}.",
        )
        db.session.commit()
        flash("Gathering form locked. Customer can no longer update.", "success")
        return redirect(url_for("crm.opportunity_edit", customer_id=gr.customer_id, tab="gathering"))

    @bp.route("/gathering/request/<int:request_id>/unlock", methods=["POST"])
    @login_required
    def gathering_request_unlock(request_id):
        gr = GatheringRequest.query.get_or_404(request_id)
        gr.is_locked = False
        gr.locked_at = None
        gr.locked_by = None
        log_opportunity_history(
            gr.customer_id,
            action="gathering-unlocked",
            changes_summary=f"Gathering request #{gr.id} unlocked by {actor_name()}.",
        )
        db.session.commit()
        flash("Gathering form unlocked. Customer can update again.", "success")
        return redirect(url_for("crm.opportunity_edit", customer_id=gr.customer_id, tab="gathering"))

    @bp.route("/gathering/request/<int:request_id>/delete", methods=["POST"])
    @login_required
    def gathering_request_delete(request_id):
        """Delete a gathering link while preserving collected data rows."""
        gr = GatheringRequest.query.get_or_404(request_id)
        customer_id = gr.customer_id

        # Keep child rows on customer by unlinking from this request.
        GatheringServerDetail.query.filter_by(gathering_request_id=request_id).update(
            {"gathering_request_id": None}, synchronize_session=False
        )
        GatheringFileNasDetail.query.filter_by(gathering_request_id=request_id).update(
            {"gathering_request_id": None}, synchronize_session=False
        )
        GatheringBlockStorageDetail.query.filter_by(gathering_request_id=request_id).update(
            {"gathering_request_id": None}, synchronize_session=False
        )

        log_opportunity_history(
            customer_id,
            action="gathering-link-deleted",
            changes_summary=(
                f"Gathering link #{request_id} (token ...{gr.token[-6:]}) deleted by {actor_name()}. "
                "Collected data preserved."
            ),
        )
        db.session.delete(gr)
        db.session.commit()
        flash("Gathering link deleted. All previously collected data has been preserved.", "success")
        return redirect(url_for("crm.opportunity_edit", customer_id=customer_id, tab="gathering"))

    @bp.route("/gathering/request/<int:request_id>/renew", methods=["POST"])
    @login_required
    def gathering_request_renew(request_id):
        """Generate a brand-new link for the same customer while keeping old data untouched."""
        gr = GatheringRequest.query.get_or_404(request_id)
        customer = Customer.query.get_or_404(gr.customer_id)
        expiry_days = request.form.get("expiry_days", type=int) or 7

        new_token = secrets.token_urlsafe(24)
        new_access_key = generate_access_key()
        new_gr = GatheringRequest(
            token=new_token,
            customer_id=customer.id,
            note=gr.note,
            expires_at=datetime.utcnow() + timedelta(days=expiry_days),
            access_key_hash=generate_password_hash(new_access_key),
            access_key_hint=new_access_key[-4:],
            status="sent",
        )
        db.session.add(new_gr)
        log_opportunity_history(
            customer.id,
            action="gathering-link-renewed",
            changes_summary=(
                f"New gathering link generated by {actor_name()} (replaces #{request_id}). "
                f"Expiry: {expiry_days} days."
            ),
        )
        db.session.commit()
        expires_text = new_gr.expires_at.strftime("%Y-%m-%d %H:%M UTC")
        flash(
            f"New link generated - copy the access key now (it will not be shown again): "
            f"{new_access_key} | Expires: {expires_text}",
            "success",
        )
        return redirect(url_for("crm.opportunity_edit", customer_id=customer.id, tab="gathering"))

    @bp.route("/gathering/form/<token>/sample.xlsx")
    def gathering_form_sample_xlsx(token):
        GatheringRequest.query.filter_by(token=token).first_or_404()

        try:
            import openpyxl  # noqa: PLC0415
        except ImportError:
            return "openpyxl is not installed on server.", 500

        wb = openpyxl.Workbook()
        ws_instructions = wb.active
        ws_instructions.title = "Instructions"
        ws_instructions.append([
            "This workbook is the data import template for AWS migration assessment. Fill the applicable sheets and upload from the customer form."
        ])

        ws_glossary = wb.create_sheet("Glossary")
        ws_glossary.append(["Attribute Name", "Example", "Requirement", "Notes"])
        ws_glossary.append(["Server Name", "Apache01", "Required for Template sheet", "Unique workload identifier"])
        ws_glossary.append(["File Server/Share Name", "CorpFiles01", "Required for FileNAS rows", "Used for NAS import mapping"])
        ws_glossary.append(["Volume Name", "vol-oracle-prod", "Required for Block Storage rows", "Used for block import mapping"])

        ws_server = wb.create_sheet("Template")
        ws_server.append([
            "Server Name", "CPU Cores", "Memory (MB)", "Provisioned Storage (GB)",
            "Operating System", "Is Virtual?", "Hypervisor Name", "Cpu String",
            "Environment", "SQL Edition", "Application",
            "Cpu Utilization Peak (%)", "Memory Utilization Peak (%)", "Time In-Use (%)",
            "Annual Cost (USD)", "Storage Type",
        ])
        ws_server.append([
            "Apache01", 4, 4096, 500,
            "Windows Server 2012 R2", "Yes", "Host-1", "Intel Xeon E7-8893 v4 @ 3.2GHz",
            "Production", "SQL Server 2012 Enterprise", "Service Now",
            0.60, 0.95, 1.00,
            3400, "HDD",
        ])

        ws_file = wb.create_sheet("FileNAS Storage (If applicable)")
        ws_file.append([
            "File Server/Share Name", "Total Used Capacity (GB) - Usable", "Access Protocol (CIFS/NFS)",
            "Total Provisioned Capacity (GB) - Usable", "Storage Efficiency Ratio (Dedupe/Compression)",
            "Peak  IOPS (reads & writes)", "Peak Throughput (MBps)", "Average  IOPS (reads & writes)",
            "Average Throughput (MBps)", "Storage Pool Name", "Array Name", "Array Vendor",
            "Average Latency (ms)", "Application",
        ])
        ws_file.append([
            "CorpFiles01", 820, "CIFS", 1200, 1.35,
            14500, 1200, 9100, 760, "Pool-A", "NetApp01", "NetApp", 1.8, "File Sharing",
        ])

        ws_block = wb.create_sheet("Block Storage (If applicable)")
        ws_block.append([
            "Volume Name", "Total Used Capacity (GB) - Usable", "Total Provisioned Capacity (GB) - Usable",
            "Peak  IOPS (reads & writes)", "Peak Throughput (MBps)", "Average  IOPS (reads & writes)",
            "Average Throughput (MBps)", "Array Name", "Average Latency (ms)", "Application",
        ])
        ws_block.append([
            "vol-oracle-prod", 2600, 3200, 42000, 1800, 30000, 1260, "PureArray01", 1.2, "Oracle DB",
        ])

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="gathering_data_sample.xlsx",
        )

    @bp.route("/gathering/form/<token>/sample-server.csv")
    def gathering_form_sample_server_csv(token):
        GatheringRequest.query.filter_by(token=token).first_or_404()

        columns = [
            "Server Name", "CPU Cores", "Memory (MB)", "Provisioned Storage (GB)",
            "Operating System", "Is Virtual (Yes/No)", "Hypervisor Name", "CPU String",
            "Environment", "SQL Edition", "Application",
            "CPU Util Peak (%)", "Memory Util Peak (%)", "Time In-Use (%)",
            "Annual Cost (USD)", "Storage Type",
        ]
        sample = [
            "Apache01", "4", "4096", "500",
            "Windows Server 2012 R2", "Yes", "Host-1", "Intel Xeon E7-8893 v4 @ 3.2GHz",
            "Production", "SQL Server 2012 Enterprise", "Service Now",
            "0.60", "0.95", "1.00",
            "3400", "HDD",
        ]

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(columns)
        writer.writerow(sample)
        output.seek(0)
        return send_file(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            mimetype="text/csv",
            as_attachment=True,
            download_name="gathering_data_sample.csv",
        )

    @bp.route("/gathering/form/<token>/sample-file-nas.csv")
    def gathering_form_sample_file_nas_csv(token):
        GatheringRequest.query.filter_by(token=token).first_or_404()

        columns = [
            "File Server/Share Name",
            "Total Used Capacity (GB) - Usable",
            "Access Protocol (CIFS/NFS)",
            "Total Provisioned Capacity (GB) - Usable",
            "Storage Efficiency Ratio (Dedupe/Compression)",
            "Peak  IOPS (reads & writes)",
            "Peak Throughput (MBps)",
            "Average  IOPS (reads & writes)",
            "Average Throughput (MBps)",
            "Storage Pool Name",
            "Array Name",
            "Array Vendor",
            "Average Latency (ms)",
            "Application",
        ]
        sample = [
            "CorpFiles01", "820", "CIFS", "1200", "1.35", "14500", "1200", "9100", "760",
            "Pool-A", "NetApp01", "NetApp", "1.8", "File Sharing",
        ]

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(columns)
        writer.writerow(sample)
        output.seek(0)
        return send_file(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            mimetype="text/csv",
            as_attachment=True,
            download_name="gathering_file_nas_sample.csv",
        )

    @bp.route("/gathering/form/<token>/sample-block-storage.csv")
    def gathering_form_sample_block_storage_csv(token):
        GatheringRequest.query.filter_by(token=token).first_or_404()

        columns = [
            "Volume Name",
            "Total Used Capacity (GB) - Usable",
            "Total Provisioned Capacity (GB) - Usable",
            "Peak  IOPS (reads & writes)",
            "Peak Throughput (MBps)",
            "Average  IOPS (reads & writes)",
            "Average Throughput (MBps)",
            "Array Name",
            "Average Latency (ms)",
            "Application",
        ]
        sample = [
            "vol-oracle-prod", "2600", "3200", "42000", "1800", "30000", "1260", "PureArray01", "1.2", "Oracle DB",
        ]

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(columns)
        writer.writerow(sample)
        output.seek(0)
        return send_file(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            mimetype="text/csv",
            as_attachment=True,
            download_name="gathering_block_storage_sample.csv",
        )
