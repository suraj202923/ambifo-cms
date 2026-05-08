from datetime import datetime, timedelta
import base64
import csv
import html
import io
import json
import os
from pathlib import Path
import re
import secrets
import signal
import threading
import urllib.request
import uuid

try:
    import dns.resolver
except Exception:
    dns = None

from flask import (
    Blueprint,
    current_app,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from flask_login import current_user, login_required, login_user, logout_user
from werkzeug.datastructures import FileStorage
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

from app.models import (
    Customer,
    CustomerDiagram,
    CustomerDocument,
    CustomerSOW,
    EmailBulkCsvExecution,
    EmailLog,
    EmailUnsubscribe,
    EmailTemplate,
    GatheringBlockStorageDetail,
    GatheringFileNasDetail,
    GatheringRequest,
    GatheringServerDetail,
    Lead,
    MeetingAvailabilityRequest,
    MeetingInvite,
    OpportunityHistory,
    OpportunityUpdateTag,
    OpportunityCloudOperator,
    OpportunityFinancial,
    OpportunitySegment,
    OpportunityStatus,
    PartnerReferenceActivity,
    PartnerReferenceContact,
    PartnerReferenceOpportunity,
    SOWMasterTemplate,
    SOWMasterTemplateSection,
    SystemSetting,
    User,
    db,
)
from app.services.email_service import EmailService
from app.services.ai_service import AIService
from app.services.macro_service import build_macro_values, render_macros
from app.configuration_admin import handle_configuration_admin_action
from app.configuration_storage import handle_configuration_storage_action
from app.gathering_import_utils import (
    _detect_rows_type,
    _import_block_storage_rows,
    _import_file_nas_rows,
    _import_server_rows,
    _import_uploaded_file_for_target,
    _import_uploaded_file_to_all_tables,
    _read_csv_rows,
    _read_xlsx_rows_by_sheet,
)
from app.routes_ai import register_ai_routes
from app.routes_auth import register_auth_routes
from app.routes_dashboard import register_dashboard_routes
from app.routes_email import register_email_routes
from app.routes_engagement import register_engagement_routes
from app.routes_gathering_data import register_gathering_data_routes
from app.routes_gathering_requests import register_gathering_request_routes
from app.routes_leads import register_lead_routes
from app.routes_opportunity_core import register_opportunity_core_routes
from app.routes_opportunity_diagrams import register_opportunity_diagram_routes
from app.routes_opportunity_documents import register_opportunity_document_routes
from app.routes_references import register_reference_routes
from app.routes_templates import register_template_routes
from app.services.admin_notification_service import (
    get_notification_tag_names,
    load_admin_history_notification_subscriptions,
    send_admin_history_notifications,
)
from app.services.storage_service import StorageService
from app.services.teams_service import TeamsService


crm_bp = Blueprint("crm", __name__)

DEFAULT_STATUSES = [
    ("In Pipeline", "#f4a259"),
    ("Qualified", "#2a9d8f"),
    ("Proposal", "#4d96ff"),
    ("Won", "#1a936f"),
    ("Lost", "#d1495b"),
]

DEFAULT_SEGMENTS = ["SUP", "SMB", "Enterprise", "Startup"]
DEFAULT_CLOUD_OPERATORS = ["Azure", "AWS", "GCP", "Other"]

SYSTEM_EMAIL_TEMPLATE_DEFAULTS = {
    "smtp_test": {
        "name": "System - SMTP Test Email",
        "subject": "SMTP Test Email - Ambifo CRM",
        "body": (
            "<p>This is a test email from Ambifo CRM SMTP Configuration.</p>"
            "<p>If you received this, SMTP is working correctly.</p>"
        ),
    },
    "gathering_request": {
        "name": "System - Gathering Information Request",
        "subject": "Ambifo - Gathering Sheet & Information Request",
        "body": (
            "<p>Hello {{customer_name}},</p>"
            "<p>Please share your required details using the secure form link below:</p>"
            "<p><a href='{{gathering_form_link}}'>Open Gathering Form</a></p>"
            "<p><strong>Access Key:</strong> {{gathering_access_key}}</p>"
            "<p><strong>Link Expiry:</strong> {{gathering_expires_at}}</p>"
            "<p>If needed, you can also upload your gathering sheet directly in this form.</p>"
            "<p>{{note}}</p>"
            "<p>Regards,<br>Ambifo Team</p>"
        ),
    },
    "meeting_link": {
        "name": "System - Meeting Invite",
        "subject": "{{meeting_subject}}",
        "body": (
            "<p>Hello {{customer_name}},</p>"
            "<p>Please find the meeting details below:</p>"
            "<p><strong>Subject:</strong> {{meeting_subject}}</p>"
            "<p><strong>When:</strong> {{meeting_when}}</p>"
            "<p><strong>Agenda:</strong><br>{{meeting_agenda_html}}</p>"
            "<p><strong>Required Data:</strong><br>{{meeting_required_data_html}}</p>"
            "<p><strong>Join Microsoft Teams:</strong><br><a href='{{meeting_link}}'>{{meeting_link_html}}</a></p>"
            "<p>Regards,<br>Ambifo Team</p>"
        ),
    },
    "meeting_availability": {
        "name": "System - Meeting Availability Request",
        "subject": "{{meeting_subject}}",
        "body": (
            "<p>Hello {{customer_name}},</p>"
            "<p>Please select your preferred meeting time by clicking one option below:</p>"
            "<p><a href='{{option_1_link}}' style='display:inline-block;padding:10px 16px;background:#1565c0;color:#fff;text-decoration:none;border-radius:8px;'>Option 1: {{option_1_text}}</a></p>"
            "<p><a href='{{option_2_link}}' style='display:inline-block;padding:10px 16px;background:#2e7d32;color:#fff;text-decoration:none;border-radius:8px;'>Option 2: {{option_2_text}}</a></p>"
            "<p><a href='{{option_3_link}}' style='display:inline-block;padding:10px 16px;background:#ef6c00;color:#fff;text-decoration:none;border-radius:8px;'>Option 3: {{option_3_text}}</a></p>"
            "<p>If these times don't work, share your own 3 preferred slots here: <a href='{{meeting_availability_form_link}}'>Open Availability Form</a></p>"
            "<p><small>This link expires on {{meeting_availability_expires_at}}.</small></p>"
            "<p>Regards,<br>Ambifo Team</p>"
        ),
    },
    "sow_document_email": {
        "name": "System - SOW Document Email",
        "subject": "{{sow_email_subject}}",
        "body": (
            "<p>{{sow_email_body_html}}</p>"
            "<p style='margin-top:16px;font-size:12px;color:#666;'>"
            "Please find the Statement of Work (DOCX) attached.</p>"
        ),
    },
}


def _system_template_name_set():
    return {cfg.get("name") for cfg in SYSTEM_EMAIL_TEMPLATE_DEFAULTS.values() if cfg.get("name")}


def _is_system_template_name(name):
    return (name or "") in _system_template_name_set()


def _ensure_default_statuses():
    if OpportunityStatus.query.count() > 0:
        return

    for name, color in DEFAULT_STATUSES:
        db.session.add(OpportunityStatus(name=name, color=color, is_active=True))
    db.session.commit()


def _ensure_default_segments():
    if OpportunitySegment.query.count() > 0:
        return

    for name in DEFAULT_SEGMENTS:
        db.session.add(OpportunitySegment(name=name, is_active=True))
    db.session.commit()


def _ensure_default_cloud_operators():
    if OpportunityCloudOperator.query.count() > 0:
        return

    for name in DEFAULT_CLOUD_OPERATORS:
        db.session.add(OpportunityCloudOperator(name=name, is_active=True))
    db.session.commit()


def _status_color_map(statuses):
    return {status.name: status.color for status in statuses}


def _user_color_map(users):
    palette = [
        "#3b82f6",
        "#0f766e",
        "#a16207",
        "#7c3aed",
        "#be123c",
        "#0891b2",
        "#6d28d9",
        "#1d4ed8",
    ]
    mapped = {}
    for idx, user in enumerate(users):
        mapped[user.id] = palette[idx % len(palette)]
    return mapped


def _get_smtp_settings():
    return {
        "host": SystemSetting.get_value("smtp.host", current_app.config.get("SMTP_HOST") or ""),
        "port": SystemSetting.get_value("smtp.port", str(current_app.config.get("SMTP_PORT") or "587")),
        "username": SystemSetting.get_value("smtp.username", current_app.config.get("SMTP_USERNAME") or ""),
        "password": SystemSetting.get_value("smtp.password", current_app.config.get("SMTP_PASSWORD") or ""),
        "use_tls": str(
            SystemSetting.get_value("smtp.use_tls", str(current_app.config.get("SMTP_USE_TLS", True)))
        ).lower()
        in ("true", "1", "yes", "on"),
        "mail_from": SystemSetting.get_value("smtp.mail_from", current_app.config.get("MAIL_FROM") or ""),
    }


def _get_teams_settings():
    return {
        "tenant_id": SystemSetting.get_value("teams.tenant_id", current_app.config.get("TEAMS_TENANT_ID") or ""),
        "client_id": SystemSetting.get_value("teams.client_id", current_app.config.get("TEAMS_CLIENT_ID") or ""),
        "client_secret": SystemSetting.get_value("teams.client_secret", current_app.config.get("TEAMS_CLIENT_SECRET") or ""),
        "organizer_id": SystemSetting.get_value("teams.organizer_id", current_app.config.get("TEAMS_ORGANIZER_ID") or ""),
        "default_duration_minutes": SystemSetting.get_value(
            "teams.default_duration_minutes",
            str(current_app.config.get("TEAMS_DEFAULT_DURATION_MINUTES") or "60"),
        ),
    }


def _get_meeting_settings():
    raw_days = (SystemSetting.get_value("meeting.availability_link_expiry_days", "7") or "7").strip()
    try:
        expiry_days = int(raw_days)
    except ValueError:
        expiry_days = 7
    expiry_days = max(1, min(90, expiry_days))
    return {
        "availability_link_expiry_days": str(expiry_days),
    }


def _get_effective_app_base_url():
    configured = (SystemSetting.get_value("app.base_url", "") or "").strip()
    if configured:
        return configured.rstrip("/")
    return (current_app.config.get("APP_BASE_URL") or "http://127.0.0.1:5000").rstrip("/")


def _appsettings_path():
    return Path(current_app.root_path).parent / "appsettings.json"


def _load_appsettings_json():
    path = _appsettings_path()
    if not path.exists():
        return {}
    try:
        with path.open("r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except (OSError, json.JSONDecodeError):
        return {}


def _save_appsettings_json(data):
    path = _appsettings_path()
    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
        f.write("\n")


def _get_document_storage_settings():
    data = _load_appsettings_json()
    storage = data.get("Storage") if isinstance(data.get("Storage"), dict) else {}
    azure = storage.get("Azure") if isinstance(storage.get("Azure"), dict) else {}
    aws = storage.get("AWS") if isinstance(storage.get("AWS"), dict) else {}
    gcp = storage.get("GCP") if isinstance(storage.get("GCP"), dict) else {}

    backend = (storage.get("Backend") or current_app.config.get("DOCUMENT_STORAGE") or "local").strip().lower()
    if backend not in {"local", "azure", "aws", "gcp"}:
        backend = "local"

    return {
        "backend": backend,
        "documents_folder": str(storage.get("DocumentsFolder") or current_app.config.get("DOCUMENTS_FOLDER") or "documents"),
        "azure_connection_string": str(azure.get("ConnectionString") or current_app.config.get("AZURE_CONNECTION_STRING") or ""),
        "azure_container_name": str(azure.get("ContainerName") or current_app.config.get("AZURE_CONTAINER_NAME") or "ambifo-documents"),
        "aws_access_key_id": str(aws.get("AccessKeyId") or current_app.config.get("AWS_ACCESS_KEY_ID") or ""),
        "aws_secret_access_key": str(aws.get("SecretAccessKey") or current_app.config.get("AWS_SECRET_ACCESS_KEY") or ""),
        "aws_bucket_name": str(aws.get("BucketName") or current_app.config.get("AWS_S3_BUCKET") or ""),
        "aws_region": str(aws.get("Region") or current_app.config.get("AWS_S3_REGION") or "us-east-1"),
        "gcp_project_id": str(gcp.get("ProjectId") or current_app.config.get("GCP_PROJECT_ID") or ""),
        "gcp_bucket_name": str(gcp.get("BucketName") or current_app.config.get("GCP_BUCKET_NAME") or ""),
        "gcp_credentials_json": str(gcp.get("CredentialsJson") or current_app.config.get("GCP_CREDENTIALS_JSON") or ""),
    }


def _apply_document_storage_runtime(settings):
    backend = (settings.get("backend") or "local").strip().lower()
    current_app.config["DOCUMENT_STORAGE"] = backend

    docs_folder = (settings.get("documents_folder") or "documents").strip()
    docs_path = Path(docs_folder)
    if not docs_path.is_absolute():
        docs_path = Path(current_app.root_path).parent / docs_path
    docs_path.mkdir(parents=True, exist_ok=True)
    current_app.config["DOCUMENTS_FOLDER"] = str(docs_path)

    current_app.config["AZURE_CONNECTION_STRING"] = (settings.get("azure_connection_string") or "").strip()
    current_app.config["AZURE_CONTAINER_NAME"] = (settings.get("azure_container_name") or "ambifo-documents").strip()

    current_app.config["AWS_ACCESS_KEY_ID"] = (settings.get("aws_access_key_id") or "").strip()
    current_app.config["AWS_SECRET_ACCESS_KEY"] = (settings.get("aws_secret_access_key") or "").strip()
    current_app.config["AWS_S3_BUCKET"] = (settings.get("aws_bucket_name") or "").strip()
    current_app.config["AWS_S3_REGION"] = (settings.get("aws_region") or "us-east-1").strip()

    current_app.config["GCP_PROJECT_ID"] = (settings.get("gcp_project_id") or "").strip()
    current_app.config["GCP_BUCKET_NAME"] = (settings.get("gcp_bucket_name") or "").strip()
    current_app.config["GCP_CREDENTIALS_JSON"] = (settings.get("gcp_credentials_json") or "").strip()


def _persist_document_storage_settings(settings):
    data = _load_appsettings_json()
    data.setdefault("Storage", {})
    data["Storage"]["Backend"] = (settings.get("backend") or "local").strip().lower()
    data["Storage"]["DocumentsFolder"] = (settings.get("documents_folder") or "documents").strip()
    data["Storage"].setdefault("Azure", {})
    data["Storage"]["Azure"]["ConnectionString"] = (settings.get("azure_connection_string") or "").strip()
    data["Storage"]["Azure"]["ContainerName"] = (settings.get("azure_container_name") or "ambifo-documents").strip()
    data["Storage"].setdefault("AWS", {})
    data["Storage"]["AWS"]["AccessKeyId"] = (settings.get("aws_access_key_id") or "").strip()
    data["Storage"]["AWS"]["SecretAccessKey"] = (settings.get("aws_secret_access_key") or "").strip()
    data["Storage"]["AWS"]["BucketName"] = (settings.get("aws_bucket_name") or "").strip()
    data["Storage"]["AWS"]["Region"] = (settings.get("aws_region") or "us-east-1").strip()
    data["Storage"].setdefault("GCP", {})
    data["Storage"]["GCP"]["ProjectId"] = (settings.get("gcp_project_id") or "").strip()
    data["Storage"]["GCP"]["BucketName"] = (settings.get("gcp_bucket_name") or "").strip()
    data["Storage"]["GCP"]["CredentialsJson"] = (settings.get("gcp_credentials_json") or "").strip()
    _save_appsettings_json(data)


def _ensure_system_email_template(template_key):
    cfg = SYSTEM_EMAIL_TEMPLATE_DEFAULTS.get(template_key)
    if not cfg:
        return None
    tpl = EmailTemplate.query.filter_by(name=cfg["name"]).first()
    if tpl:
        return tpl
    tpl = EmailTemplate(
        name=cfg["name"],
        subject_template=cfg["subject"],
        body_template=cfg["body"],
    )
    db.session.add(tpl)
    db.session.flush()
    return tpl


def _render_system_email_template(template_key, macro_values):
    tpl = _ensure_system_email_template(template_key)
    if not tpl:
        return None, "", ""
    rendered_subject = render_macros(tpl.subject_template or "", macro_values).strip()
    rendered_body = render_macros(tpl.body_template or "", macro_values).strip()
    return tpl, rendered_subject, rendered_body


def ensure_default_system_email_templates():
    ensured_any = False
    for template_key in SYSTEM_EMAIL_TEMPLATE_DEFAULTS.keys():
        tpl = _ensure_system_email_template(template_key)
        if tpl is not None and tpl.id is not None:
            ensured_any = True
    return ensured_any


def _get_ai_settings():
    active_provider = (
        SystemSetting.get_value("ai.active_provider", None)
        or SystemSetting.get_value("ai.provider", "openai")
        or "openai"
    ).strip().lower()

    return {
        "active_provider": active_provider,
        "openai_api_key": (
            SystemSetting.get_value("ai.openai.api_key", None)
            or SystemSetting.get_value("ai.api_key", "")
            or ""
        ).strip(),
        "openai_model": (SystemSetting.get_value("ai.openai.model", "gpt-4o-mini") or "gpt-4o-mini").strip(),
        "openai_base_url": (SystemSetting.get_value("ai.openai.base_url", "https://api.openai.com/v1") or "https://api.openai.com/v1").strip(),
        "gemini_api_key": (SystemSetting.get_value("ai.gemini.api_key", "") or "").strip(),
        "gemini_model": (SystemSetting.get_value("ai.gemini.model", "gemini-1.5-flash") or "gemini-1.5-flash").strip(),
        "gemini_base_url": (SystemSetting.get_value("ai.gemini.base_url", "https://generativelanguage.googleapis.com/v1beta") or "https://generativelanguage.googleapis.com/v1beta").strip(),
    }


def _build_ai_customer_context(customer):
    server_rows = GatheringServerDetail.query.filter_by(customer_id=customer.id).order_by(GatheringServerDetail.created_at.desc()).limit(12).all()
    file_rows = GatheringFileNasDetail.query.filter_by(customer_id=customer.id).order_by(GatheringFileNasDetail.created_at.desc()).limit(12).all()
    block_rows = GatheringBlockStorageDetail.query.filter_by(customer_id=customer.id).order_by(GatheringBlockStorageDetail.created_at.desc()).limit(12).all()

    lines = [
        f"Customer Name: {customer.customer_name or ''}",
        f"Account Name: {customer.account_name or ''}",
        f"Cloud: {customer.cloud or ''}",
        f"Segment: {customer.segment or ''}",
        f"Status: {customer.deal_status or ''}",
        f"Comment: {customer.comment or ''}",
        f"Next Action: {customer.next_action_planned or ''}",
        "",
        "Server Data:",
    ]
    for row in server_rows:
        lines.append(
            f"- {row.server_name or 'server'} | CPU:{row.cpu_cores or ''} | MemoryMB:{row.memory_mb or ''} | "
            f"StorageGB:{row.provisioned_storage_gb or ''} | OS:{row.operating_system or ''} | App:{row.application or ''}"
        )

    lines.append("")
    lines.append("File/NAS Data:")
    for row in file_rows:
        lines.append(
            f"- {row.file_server_share_name or 'share'} | UsedGB:{row.total_used_capacity_gb or ''} | "
            f"ProvisionedGB:{row.total_provisioned_capacity_gb or ''} | Protocol:{row.access_protocol or ''} | App:{row.application or ''}"
        )

    lines.append("")
    lines.append("Block Storage Data:")
    for row in block_rows:
        lines.append(
            f"- {row.volume_name or 'volume'} | UsedGB:{row.total_used_capacity_gb or ''} | "
            f"ProvisionedGB:{row.total_provisioned_capacity_gb or ''} | Array:{row.array_name or ''} | App:{row.application or ''}"
        )

    return "\n".join(lines)


def _actor_name():
    if current_user and current_user.is_authenticated:
        return current_user.username
    return "system"


def _tag_color_map():
    tags = OpportunityUpdateTag.query.filter_by(is_active=True).all()
    return {t.name.lower(): t.color for t in tags}


def _derive_tag_from_action(action):
    action_text = (action or "").lower()
    if "document" in action_text:
        return "document"
    if "meeting" in action_text:
        return "meeting"
    if action_text.startswith("email"):
        return "email"
    if "gathering" in action_text:
        return "gathering"
    if "status" in action_text:
        return "status"
    if "segment" in action_text:
        return "segment"
    if "assign" in action_text:
        return "assign"
    return "update"


def _log_opportunity_history(customer_id, action, changes_summary, remark=None, tag_name=None):
    final_tag = (tag_name or "").strip().lower() or _derive_tag_from_action(action)
    db.session.add(
        OpportunityHistory(
            customer_id=customer_id,
            changed_by=_actor_name(),
            action=action,
            tag_name=final_tag,
            changes_summary=changes_summary,
            remark=remark,
        )
    )
    try:
        send_admin_history_notifications(
            customer_id,
            action,
            changes_summary,
            final_tag,
            changed_by=_actor_name(),
        )
    except Exception:
        # Notification delivery failures should not block primary business flow.
        pass


def _to_int(value):
    raw = (value or "").strip()
    if not raw:
        return None
    return int(raw)


def _to_float(value):
    raw = (value or "").strip()
    if not raw:
        return None
    return float(raw)


def _to_optional_float(value, label):
    raw = str(value).strip() if value is not None else ""
    if not raw:
        return None
    try:
        return float(raw)
    except (TypeError, ValueError):
        raise ValueError(f"{label} must be a valid number.")


def _round_money(value):
    if value is None:
        return None
    return round(float(value), 2)


def _resolve_arr_value(mrr_value, arr_value):
    if arr_value is not None:
        return _round_money(arr_value)
    if mrr_value is None:
        return None
    return _round_money(mrr_value * 12)


def _compute_credit_value(base_value, credit_percentage):
    if base_value is None or credit_percentage is None:
        return None
    return _round_money((base_value * credit_percentage) / 100.0)


def _segment_credit_rule(segment_name):
    name = (segment_name or "").strip()
    if not name:
        return None, False
    seg = OpportunitySegment.query.filter_by(name=name).first()
    if not seg:
        return None, None, "mrr"

    pct_customer = seg.credit_percentage_customer
    pct_ambifo = seg.credit_percentage_ambifo

    if pct_customer is None and seg.credit_percentage is not None:
        pct_customer = seg.credit_percentage
    if pct_ambifo is None and seg.credit_percentage is not None:
        pct_ambifo = seg.credit_percentage

    basis = "arr" if bool(seg.credit_basis_arr) else "mrr"
    if bool(seg.credit_basis_arr) and bool(seg.credit_basis_mrr):
        basis = "arr"
    elif not bool(seg.credit_basis_arr) and not bool(seg.credit_basis_mrr):
        basis = "mrr"

    return pct_customer, pct_ambifo, basis


def _build_financial_payload(form, segment_name):
    expected_mrr = _to_optional_float(form.get("expected_mrr"), "Expected MRR")
    expected_arr_input = _to_optional_float(form.get("expected_arr"), "Expected ARR")
    actual_mrr = _to_optional_float(form.get("actual_mrr"), "Actual MRR")
    actual_arr_input = _to_optional_float(form.get("actual_arr"), "Actual ARR")

    expected_arr = _resolve_arr_value(expected_mrr, expected_arr_input)
    actual_arr = _resolve_arr_value(actual_mrr, actual_arr_input)

    credit_pct_customer, credit_pct_ambifo, basis = _segment_credit_rule(segment_name)
    expected_credit_base = expected_arr if basis == "arr" else expected_mrr
    actual_credit_base = actual_arr if basis == "arr" else actual_mrr

    expected_credit_customer = _compute_credit_value(expected_credit_base, credit_pct_customer)
    expected_credit_ambifo = _compute_credit_value(expected_credit_base, credit_pct_ambifo)
    actual_credit_customer = _compute_credit_value(actual_credit_base, credit_pct_customer)
    actual_credit_ambifo = _compute_credit_value(actual_credit_base, credit_pct_ambifo)

    return {
        "expected_mrr": _round_money(expected_mrr),
        "expected_arr": expected_arr,
        "expected_credit_customer": expected_credit_customer,
        "expected_credit_ambifo": expected_credit_ambifo,
        "actual_mrr": _round_money(actual_mrr),
        "actual_arr": actual_arr,
        "actual_credit_customer": actual_credit_customer,
        "actual_credit_ambifo": actual_credit_ambifo,
    }


def _financial_change_lines(existing_financial, payload):
    labels = {
        "expected_mrr": "Expected MRR",
        "expected_arr": "Expected ARR",
        "expected_credit_customer": "Expected Credit To Customer",
        "expected_credit_ambifo": "Expected Credit To Ambifo",
        "actual_mrr": "Actual MRR",
        "actual_arr": "Actual ARR",
        "actual_credit_customer": "Actual Credit To Customer",
        "actual_credit_ambifo": "Actual Credit To Ambifo",
    }
    lines = []
    for key, label in labels.items():
        old_value = getattr(existing_financial, key, None) if existing_financial else None
        new_value = payload.get(key)
        if old_value is None and new_value is None:
            continue
        if _round_money(old_value) != _round_money(new_value):
            lines.append(f"{label}: '{'' if old_value is None else _round_money(old_value)}' -> '{'' if new_value is None else _round_money(new_value)}'")
    return lines


def _apply_financial_payload(customer_id, payload):
    financial = OpportunityFinancial.query.filter_by(customer_id=customer_id).first()
    if not financial:
        financial = OpportunityFinancial(customer_id=customer_id)
        db.session.add(financial)

    for key, value in payload.items():
        setattr(financial, key, value)

    return financial


def _recalculate_financials_for_segment(segment_name):
    customers = Customer.query.filter(Customer.segment == (segment_name or "").strip()).all()
    for customer in customers:
        financial = customer.financial_profile
        if not financial:
            continue
        payload = _build_financial_payload(
            {
                "expected_mrr": financial.expected_mrr,
                "expected_arr": financial.expected_arr,
                "actual_mrr": financial.actual_mrr,
                "actual_arr": financial.actual_arr,
            },
            customer.segment,
        )
        _apply_financial_payload(customer.id, payload)


def _generate_access_key(length=8):
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
    return "".join(secrets.choice(alphabet) for _ in range(length))


def _is_gathering_request_expired(request_record):
    return bool(request_record.expires_at and datetime.utcnow() > request_record.expires_at)


def _request_process_restart(delay_seconds=1.0):
    """Restart the full service process (Gunicorn master on Render)."""
    master_pid = os.getppid()
    target_pid = master_pid if master_pid and master_pid > 1 else os.getpid()
    threading.Timer(delay_seconds, lambda: os.kill(target_pid, signal.SIGTERM)).start()




def _format_option_dt(value):
    return value.strftime("%d %b %Y, %I:%M %p UTC")


def _build_template_macro_values_for_customer(customer, template, recipient_email=None):
    macro_values = build_macro_values(customer)

    active_diagrams = (
        CustomerDiagram.query
        .filter_by(customer_id=customer.id, is_active=True)
        .order_by(CustomerDiagram.created_at.asc())
        .all()
    )
    if active_diagrams:
        diagram_html_blocks = []
        for diagram in active_diagrams:
            safe_name = html.escape(diagram.diagram_name or "Diagram")
            diagram_html_blocks.append(f"<h4>{safe_name}</h4>{diagram.diagram_content or ''}")
            diagram_macro = (diagram.macro_key or "").strip()
            if diagram_macro:
                macro_values[diagram_macro] = diagram.diagram_content or ""
        macro_values["selected_diagrams_html"] = "\n".join(diagram_html_blocks)
    else:
        macro_values["selected_diagrams_html"] = ""

    if not template:
        return macro_values

    template_text = f"{template.subject_template or ''}\n{template.body_template or ''}"
    requires_link = (
        "{{meeting_availability_form_link}}" in template_text
        or "{{meeting_availability_expires_at}}" in template_text
    )
    if not requires_link:
        return macro_values

    meeting_settings = _get_meeting_settings()
    expiry_days = int(meeting_settings["availability_link_expiry_days"])
    expires_at = datetime.utcnow() + timedelta(days=expiry_days)
    token = secrets.token_urlsafe(24)

    base_slot = datetime.utcnow() + timedelta(days=1)
    req = MeetingAvailabilityRequest(
        token=token,
        customer_id=customer.id,
        recipient_email=(recipient_email or customer.email or "").strip() or (customer.email or "unknown@local"),
        subject=(template.subject_template or "Select Your Preferred Meeting Time"),
        option_1_at=base_slot,
        option_2_at=base_slot + timedelta(hours=1),
        option_3_at=base_slot + timedelta(hours=2),
        expires_at=expires_at,
        status="sent",
        created_by=_actor_name(),
    )
    db.session.add(req)

    expires_text = expires_at.strftime("%d %b %Y, %I:%M %p UTC")
    _log_opportunity_history(
        customer.id,
        action="meeting-availability-link-generated",
        changes_summary=(
            "Meeting availability token link generated from template send.\n"
            f"Template: {template.name if getattr(template, 'name', None) else 'N/A'}\n"
            f"Recipient: {(recipient_email or customer.email or 'N/A')}\n"
            f"Token: ...{token[-6:]}\n"
            f"Expires: {expires_text}"
        ),
        remark=None,
        tag_name="meeting",
    )

    base_url = _get_effective_app_base_url()
    macro_values["meeting_availability_form_link"] = f"{base_url}/meeting/availability/form/{token}"
    macro_values["meeting_availability_expires_at"] = expires_text
    return macro_values


def _email_template_macro_catalog():
    # Canonical macro keys supported by email template rendering paths.
    keys = [
        "customer_name",
        "account_name",
        "email",
        "phone",
        "city",
        "segment",
        "deal_status",
        "today",
        "gathering_form_link",
        "gathering_access_key",
        "gathering_expires_at",
        "meeting_availability_form_link",
        "meeting_availability_expires_at",
        "selected_diagrams_html",
    ]
    return [f"{{{{{k}}}}}" for k in keys]


def _sanitize_email_template_body(body_html):
    text = (body_html or "").strip()
    if not text:
        return ""

    # If a full Ambifo wrapper was pasted, keep only the editable content cell.
    if (
        "support@ambifo.com" in text
        and "ambifo.com" in text
        and "www.linkedin.com/company/ambifo-technology" in text
    ):
        match = re.search(
            r'<td[^>]*padding:20px;line-height:1\.6;font-size:14px;[^>]*>([\s\S]*?)</td>',
            text,
            flags=re.IGNORECASE,
        )
        if match:
            return (match.group(1) or "").strip()

    # Remove explicit footer tags if users paste them.
    text = re.sub(r"<footer[\s\S]*?</footer>", "", text, flags=re.IGNORECASE)

    # Remove lines/blocks containing locked branding footer contact details.
    footer_tokens = [
        "support@ambifo.com",
        "https://ambifo.com",
        "www.linkedin.com/company/ambifo-technology",
        "+91 9148419502",
        "Building No 674",
    ]
    for token in footer_tokens:
        text = re.sub(
            rf"<(p|div|span|li)[^>]*>[\s\S]*?{re.escape(token)}[\s\S]*?</\1>",
            "",
            text,
            flags=re.IGNORECASE,
        )

    return text.strip()


def _is_valid_email(value):
    if not value:
        return False
    pattern = r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$"
    return re.match(pattern, value) is not None


def _parse_email_list(raw_value):
    raw = (raw_value or "").replace(";", ",").replace("\n", ",")
    values = [item.strip().lower() for item in raw.split(",") if item.strip()]
    duplicate_values = []
    unique_values = []
    seen = set()
    for item in values:
        if item in seen:
            if item not in duplicate_values:
                duplicate_values.append(item)
            continue
        seen.add(item)
        unique_values.append(item)
    invalid_values = [item for item in unique_values if not _is_valid_email(item)]
    valid_values = [item for item in unique_values if _is_valid_email(item)]
    return valid_values, invalid_values, duplicate_values


def _customer_alternate_email_list(customer):
    if not customer:
        return []
    valid_values, _, _ = _parse_email_list(customer.alternate_emails or "")
    return valid_values


def _merge_cc_lists(primary_recipients, manual_cc=None, alternate_cc=None):
    primary_set = {(item or "").strip().lower() for item in (primary_recipients or []) if item}
    merged = []
    seen = set()
    for item in (manual_cc or []) + (alternate_cc or []):
        normalized = (item or "").strip().lower()
        if not normalized or normalized in primary_set or normalized in seen:
            continue
        seen.add(normalized)
        merged.append(normalized)
    return merged


def _normalize_macro_key(value):
    key = re.sub(r"[^a-zA-Z0-9_]", "_", (value or "").strip().lower())
    key = re.sub(r"_+", "_", key).strip("_")
    return key


# Route blocks below extracted to dedicated modules; see register_* calls at bottom.


SOW_DEFAULT_SECTIONS = [
        (
                "Proposal Information",
                """
<h2>Proposal Information</h2>
<div><strong>Proposal For</strong></div>
<div>{{customer_name}}</div>
<div>{{company_name}}</div>
<div style='margin-top:8px;'><strong>Submitted By</strong></div>
<div>{{ambifo_company_name}}</div>
<div>{{ambifo_locations}}</div>
<div style='margin-top:8px;'><strong>SOW Title</strong></div>
<div>{{sow_title}}</div>
<div style='margin-top:4px;'><strong>Date</strong></div>
<div>{{sow_date}}</div>
<div style='margin-top:4px;'><strong>Version</strong></div>
<div>{{sow_version}}</div>
<table style='width:100%;border-collapse:collapse;margin-top:12px;border:1px solid #e2e8f0;'>
    <thead>
        <tr>
            <th style='text-align:left;padding:8px;border-bottom:1px solid #e2e8f0;background:#f8fafc;'>Field</th>
            <th style='text-align:left;padding:8px;border-bottom:1px solid #e2e8f0;background:#f8fafc;'>Details</th>
        </tr>
    </thead>
    <tbody>
        <tr><td style='padding:8px;border-top:1px solid #e2e8f0;'>Signed Between</td><td style='padding:8px;border-top:1px solid #e2e8f0;'>{{ambifo_company_name}} &amp; {{company_name}}</td></tr>
        <tr><td style='padding:8px;border-top:1px solid #e2e8f0;'>Customer Name</td><td style='padding:8px;border-top:1px solid #e2e8f0;'>{{customer_name}}</td></tr>
        <tr><td style='padding:8px;border-top:1px solid #e2e8f0;'>Customer Email</td><td style='padding:8px;border-top:1px solid #e2e8f0;'>{{customer_email}}</td></tr>
        <tr><td style='padding:8px;border-top:1px solid #e2e8f0;'>Customer Phone</td><td style='padding:8px;border-top:1px solid #e2e8f0;'>{{customer_phone}}</td></tr>
        <tr><td style='padding:8px;border-top:1px solid #e2e8f0;'>Cloud Platform</td><td style='padding:8px;border-top:1px solid #e2e8f0;'>{{cloud_platform}}</td></tr>
        <tr><td style='padding:8px;border-top:1px solid #e2e8f0;'>Ambifo Contact</td><td style='padding:8px;border-top:1px solid #e2e8f0;'>{{ambifo_support_email}} | {{ambifo_website}}</td></tr>
    </tbody>
</table>
                """.strip(),
        ),
    ("About Ambifo", "<p>Ambifo Technology Pvt Ltd is an AWS &amp; Azure Premier Consulting Partner, providing cloud consulting, implementation, and managed services.</p>"),
    ("Offerings", "<p>Cloud migration, managed services, DevOps, security, AI/ML and analytics services.</p>"),
    ("Business Background", "<p><strong>{{company_name}}</strong> is evaluating cloud transformation to improve scalability and resilience on <strong>{{cloud_platform}}</strong>.</p>"),
    ("Project Overview", "<p>This engagement covers migration/deployment for <strong>{{customer_name}}</strong> with implementation and support.</p>"),
    ("Problem Statement", "<p><strong>{{company_name}}</strong> currently faces scalability, cost, and operational challenges.</p>"),
    ("Document Objective", "<p>This SOW defines scope, deliverables, timelines and responsibilities.</p>"),
    ("Success Criteria", "<ul><li>Successful infrastructure provisioning</li><li>Testing and validation complete</li><li>Go-live sign-off</li></ul>"),
    ("Proposed Solution", "<p>Solution architecture, implementation approach, assumptions and out-of-scope items.</p>"),
    ("Scope & Schedule", "<p>Phased plan: discovery, provisioning, migration, testing and handover.</p>"),
    ("Project Governance", "<p>Steering, PM ownership, reporting model and escalation matrix.</p>"),
    ("Commercials / Terms / Sign-off", "<p>Pricing, terms &amp; conditions, and acceptance signatories.</p>"),
]

LEGACY_SOW_SECTION_FIELDS = [
    "section_about",
    "section_offerings",
    "section_business_background",
    "section_project_overview",
    "section_problem_statement",
    "section_document_objective",
    "section_success_criteria",
    "section_proposed_solution",
    "section_scope_schedule",
    "section_project_governance",
    "section_commercials_signoff",
]


def _get_template_sections(master_template):
    return (
        SOWMasterTemplateSection.query
        .filter_by(template_id=master_template.id)
        .order_by(SOWMasterTemplateSection.sequence_no.asc(), SOWMasterTemplateSection.id.asc())
        .all()
    )


def _ensure_template_sections(master_template):
    existing = _get_template_sections(master_template)

    if not existing:
        # First-time seed: create all default sections only when the template has none.
        created_sections = []
        for idx, (section_name, default_html) in enumerate(SOW_DEFAULT_SECTIONS, start=1):
            legacy_html = getattr(master_template, LEGACY_SOW_SECTION_FIELDS[idx - 1], None) if idx - 1 < len(LEGACY_SOW_SECTION_FIELDS) else None
            content = (legacy_html or "").strip() or default_html
            section = SOWMasterTemplateSection(
                template_id=master_template.id,
                section_name=section_name,
                sequence_no=idx,
                content_html=content,
            )
            db.session.add(section)
            created_sections.append(section)
        db.session.commit()
        return created_sections

    # Template already has sections — return them as-is without recreating deleted ones.
    return existing


def _normalize_template_section_sequence(template_id):
    sections = (
        SOWMasterTemplateSection.query
        .filter_by(template_id=template_id)
        .order_by(SOWMasterTemplateSection.sequence_no.asc(), SOWMasterTemplateSection.id.asc())
        .all()
    )
    for idx, section in enumerate(sections, start=1):
        section.sequence_no = idx


def _move_template_section_to_position(template_id, section_id, target_position):
    sections = (
        SOWMasterTemplateSection.query
        .filter_by(template_id=template_id)
        .order_by(SOWMasterTemplateSection.sequence_no.asc(), SOWMasterTemplateSection.id.asc())
        .all()
    )
    if not sections:
        return

    moving = next((s for s in sections if s.id == section_id), None)
    if not moving:
        return

    remaining = [s for s in sections if s.id != section_id]
    bounded_position = max(1, min(int(target_position or 1), len(sections)))
    insert_at = bounded_position - 1
    remaining.insert(insert_at, moving)

    for idx, section in enumerate(remaining, start=1):
        section.sequence_no = idx


def _get_sow_brand_config():
    def _int_setting(key, default):
        raw = (SystemSetting.get_value(key, str(default)) or str(default)).strip()
        try:
            value = int(raw)
        except (TypeError, ValueError):
            value = int(default)
        return max(0, min(value, 300))

    return {
        "company_name":  SystemSetting.get_value("sow.brand.company_name",  "AMBIFO TECHNOLOGY PVT LTD"),
        "tagline_chips": SystemSetting.get_value("sow.brand.tagline_chips", "Cloud Strategy,Migration & Modernization,DevOps,AI/ML"),
        "website":       SystemSetting.get_value("sow.brand.website",       "www.ambifo.com"),
        "email":         SystemSetting.get_value("sow.brand.email",         "support@ambifo.com"),
        "phone":         SystemSetting.get_value("sow.brand.phone",         "+91 9148419502"),
        "address":       SystemSetting.get_value("sow.brand.address",       "Building No 674, 18th Main, 3rd Phase, Domlur, Bangalore - 560071"),
        "confidential":  SystemSetting.get_value("sow.brand.confidential",  "⚠ CONFIDENTIAL — Intended solely for the named addressee"),
        "header_top_px": _int_setting("sow.brand.header_top_px", 122),
        "footer_bottom_px": _int_setting("sow.brand.footer_bottom_px", 80),
    }


def _get_sow_document_title_label(template_id):
    key = f"sow.template.{int(template_id)}.document_title"
    value = (SystemSetting.get_value(key, "") or "").strip()
    return value or "Statement of Work (SOW)"


def _set_sow_document_title_label(template_id, value):
    key = f"sow.template.{int(template_id)}.document_title"
    label = (value or "").strip() or "Statement of Work (SOW)"
    SystemSetting.set_value(key, label)


def _format_doc_ref_no(value=None, sow_id=None, ref_date=None):
    raw = (value or "").strip()
    if re.match(r"^AMBIFO-\d{6}-\d{3}$", raw, flags=re.IGNORECASE):
        return raw.upper()

    date_part = (ref_date or datetime.utcnow()).strftime("%d%m%y")

    if raw:
        m = re.search(r"(\d+)$", raw)
        if m:
            return f"AMBIFO-{date_part}-{int(m.group(1)):03d}"
        return raw
    if sow_id:
        try:
            return f"AMBIFO-{date_part}-{int(sow_id):03d}"
        except (TypeError, ValueError):
            pass
    return f"AMBIFO-{date_part}-000"


def _next_document_editor_ref_no(now=None):
    """Generate next shared ref no in format AMBIFO-YYMMDD-0001."""
    dt = now or datetime.utcnow()
    date_part = dt.strftime("%y%m%d")

    date_key = "doc.editor.ref.current_date"
    seq_key = "doc.editor.ref.current_seq"

    stored_date = (SystemSetting.get_value(date_key, "") or "").strip()
    raw_seq = (SystemSetting.get_value(seq_key, "0") or "0").strip()
    try:
        seq = int(raw_seq)
    except (TypeError, ValueError):
        seq = 0

    if stored_date != date_part:
        seq = 0

    seq += 1

    SystemSetting.set_value(date_key, date_part)
    SystemSetting.set_value(seq_key, str(seq))
    db.session.commit()

    return f"AMBIFO-{date_part}-{seq:04d}"


def _compose_master_template_html(master_template):
    """Build editable HTML body from SOW master template."""
    if not master_template:
        return ""

    content_html = (master_template.content_html or "").strip()
    if content_html:
        return content_html

    sections = _get_template_sections(master_template)
    chunks = []
    for section in sections:
        title = (getattr(section, "section_name", None) or getattr(section, "title", "") or "").strip()
        body = (section.content_html or "").strip()
        if title:
            chunks.append(f"<h2>{html.escape(title)}</h2>")
        if body:
            chunks.append(body)

    return "\n".join(chunks).strip()


def _create_master_template(template_name, template_key):
    master = SOWMasterTemplate(
        template_name=template_name,
        template_key=template_key,
        updated_by=_actor_name(),
    )
    for idx, legacy_field in enumerate(LEGACY_SOW_SECTION_FIELDS):
        setattr(master, legacy_field, SOW_DEFAULT_SECTIONS[idx][1])
    master.content_html = ""
    return master


def _generate_unique_sow_template_key(base_text):
    base_key = re.sub(r"[^a-z0-9]+", "-", (base_text or "").lower()).strip("-") or "template"
    key = base_key
    seq = 2
    while SOWMasterTemplate.query.filter_by(template_key=key).first():
        key = f"{base_key}-{seq}"
        seq += 1
    return key


def _ensure_default_sow_master_template():
    existing = SOWMasterTemplate.query.order_by(SOWMasterTemplate.id.asc()).all()
    if not existing:
        master = _create_master_template("Default Template", "default")
        db.session.add(master)
        db.session.commit()
        return master

    changed = False
    for idx, master in enumerate(existing, start=1):
        if not (master.template_name or "").strip():
            master.template_name = "Default Template" if idx == 1 else f"Template {idx}"
            changed = True
        for i, legacy_field in enumerate(LEGACY_SOW_SECTION_FIELDS):
            if not getattr(master, legacy_field, None):
                setattr(master, legacy_field, SOW_DEFAULT_SECTIONS[i][1])
                changed = True
        _ensure_template_sections(master)
    if changed:
        db.session.commit()
    return existing[0]


def _normalize_selected_diagram_ids(customer_id, raw_ids):
    if raw_ids is None:
        return []

    parsed = []
    if isinstance(raw_ids, list):
        for item in raw_ids:
            try:
                parsed.append(int(item))
            except (TypeError, ValueError):
                continue
    elif isinstance(raw_ids, str):
        for token in raw_ids.split(","):
            token = token.strip()
            if not token:
                continue
            try:
                parsed.append(int(token))
            except (TypeError, ValueError):
                continue
    else:
        try:
            parsed.append(int(raw_ids))
        except (TypeError, ValueError):
            pass

    if not parsed:
        return []

    valid_ids = {
        d.id for d in CustomerDiagram.query.filter(
            CustomerDiagram.customer_id == customer_id,
            CustomerDiagram.is_active == True,
            CustomerDiagram.id.in_(parsed),
        ).all()
    }
    return [item for item in parsed if item in valid_ids]


def _selected_diagrams_for_customer(customer_id, selected_diagram_ids):
    if not selected_diagram_ids:
        return (
            CustomerDiagram.query
            .filter_by(customer_id=customer_id, is_active=True)
            .order_by(CustomerDiagram.created_at.asc())
            .all()
        )

    return (
        CustomerDiagram.query
        .filter(
            CustomerDiagram.customer_id == customer_id,
            CustomerDiagram.is_active == True,
            CustomerDiagram.id.in_(selected_diagram_ids),
        )
        .order_by(CustomerDiagram.created_at.asc())
        .all()
    )


def _render_sow_template_for_customer(template_html, customer, selected_diagrams=None, sow=None):
    safe_customer_name = html.escape(customer.customer_name or "")
    safe_company_name = html.escape(customer.account_name or customer.customer_name or "")
    safe_cloud = html.escape(customer.cloud or "AWS / Azure")
    safe_email = html.escape(customer.email or "")
    safe_phone = html.escape(customer.phone or "-")
    safe_today = datetime.utcnow().strftime("%d %B, %Y")
    safe_sow_title = html.escape((getattr(sow, "sow_title", None) or f"Cloud Migration SOW — {customer.customer_name}").strip())
    safe_sow_date = html.escape((getattr(sow, "sow_date", None) or safe_today).strip())
    safe_sow_version = html.escape((getattr(sow, "version", None) or "1.0").strip())
    safe_doc_ref_no = html.escape(
        _format_doc_ref_no(
            sow_id=getattr(sow, "id", None),
            ref_date=getattr(sow, "created_at", None),
        )
    )
    template_id = getattr(sow, "master_template_id", None)
    safe_document_title_label = html.escape(
        _get_sow_document_title_label(template_id) if template_id else "Statement of Work (SOW)"
    )

    rendered = template_html or ""
    replacements = {
        "{{customer_name}}": safe_customer_name,
        "{{company_name}}": safe_company_name,
        "{{cloud_platform}}": safe_cloud,
        "{{customer_email}}": safe_email,
        "{{customer_phone}}": safe_phone,
        "{{today}}": safe_today,
        "{{sow_title}}": safe_sow_title,
        "{{sow_date}}": safe_sow_date,
        "{{sow_version}}": safe_sow_version,
        "{{doc_ref_no}}": safe_doc_ref_no,
        "{{document_title_label}}": safe_document_title_label,
        "{{ambifo_company_name}}": "Ambifo Technology Pvt Ltd",
        "{{ambifo_locations}}": "Bangalore | Mumbai | Delhi",
        "{{ambifo_website}}": "www.ambifo.com",
        "{{ambifo_support_email}}": "support@ambifo.com",
        "{{ambifo_support_phone}}": "+91 9148419502",
        "{{ambifo_address}}": "Building No 674, 18th Main, 3rd Phase, Front of New Land ISRO Quarter, Domlur, Bangalore - 560071",
        "{{ambifo_linkedin}}": "https://www.linkedin.com/company/ambifo-technology",
    }

    selected_diagrams = selected_diagrams or []
    diagram_html_blocks = []
    for diagram in selected_diagrams:
        macro_key = (diagram.macro_key or "").strip()
        if not macro_key:
            continue
        replacements[f"{{{{{macro_key}}}}}"] = diagram.diagram_content or ""
        safe_name = html.escape(diagram.diagram_name or "Diagram")
        diagram_html_blocks.append(f"<h4>{safe_name}</h4>{diagram.diagram_content or ''}")

    replacements["{{selected_diagrams_html}}"] = "\n".join(diagram_html_blocks)

    for token, value in replacements.items():
        rendered = rendered.replace(token, value)
    return rendered


def _build_rendered_sow_sections(master_template, customer, selected_diagrams=None, sow=None):
    sections = _ensure_template_sections(master_template)
    rendered_parts = []
    for idx, section in enumerate(sections, start=1):
        title = html.escape((section.section_name or "Section").strip())
        body = _render_sow_template_for_customer(
            section.content_html or "",
            customer,
            selected_diagrams=selected_diagrams,
            sow=sow,
        )
        rendered_parts.append(f"<h1>{idx}. {title}</h1>{body}")
    return "\n".join(rendered_parts)


def _append_missing_rendered_sow_sections(existing_html, master_template, customer, selected_diagrams=None, sow=None):
    sections = _ensure_template_sections(master_template)
    source_html = existing_html or ""
    missing_parts = []

    for idx, section in enumerate(sections, start=1):
        title = (section.section_name or "Section").strip()
        if not title:
            continue

        # Detect section by heading text (sequence number may vary over time).
        heading_re = re.compile(
            rf"<h1[^>]*>\s*\d+\.\s*{re.escape(title)}\s*</h1>",
            flags=re.IGNORECASE,
        )
        if heading_re.search(source_html):
            continue

        rendered_body = _render_sow_template_for_customer(
            section.content_html or "",
            customer,
            selected_diagrams=selected_diagrams,
            sow=sow,
        )
        safe_title = html.escape(title)
        missing_parts.append(f"<h1>{idx}. {safe_title}</h1>{rendered_body}")

    if not missing_parts:
        return source_html, False

    separator = "\n" if source_html.strip() else ""
    return f"{source_html}{separator}{'\n'.join(missing_parts)}", True


def _html_to_docx_body(doc, html_content):
    """Convert basic Quill HTML to python-docx paragraph elements."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html_content or "", "html.parser")

    def _add_inline(para, elem):
        for node in elem.children:
            if isinstance(node, str):
                if node:
                    para.add_run(node)
            elif node.name in ("strong", "b"):
                run = para.add_run(node.get_text())
                run.bold = True
            elif node.name in ("em", "i"):
                run = para.add_run(node.get_text())
                run.italic = True
            elif node.name == "u":
                run = para.add_run(node.get_text())
                run.underline = True
            elif node.name == "br":
                para.add_run("\n")
            elif node.name:
                para.add_run(node.get_text())

    for el in soup.children:
        if not hasattr(el, "name") or not el.name:
            if isinstance(el, str) and el.strip():
                doc.add_paragraph(el.strip())
            continue
        tag = el.name
        if tag in ("h1", "h2", "h3", "h4"):
            doc.add_heading(el.get_text(), level=int(tag[1]))
        elif tag == "p":
            text = el.get_text()
            if text.strip():
                p = doc.add_paragraph()
                _add_inline(p, el)
        elif tag in ("ul", "ol"):
            style = "List Bullet" if tag == "ul" else "List Number"
            for li in el.find_all("li", recursive=False):
                p = doc.add_paragraph(style=style)
                _add_inline(p, li)
        elif tag == "blockquote":
            from docx.shared import Pt
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(36)
            _add_inline(p, el)
        elif tag in ("div", "section"):
            _html_to_docx_body(doc, str(el.decode_contents()))


def _generate_sow_docx(customer, sow_title, sow_date, sow_version, content_html, document_title_label="Statement of Work (SOW)", doc_ref_no=""):
    """Build and return a DOCX file as bytes."""
    from io import BytesIO
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)

    BRAND_BLUE = RGBColor(0x0A, 0x1F, 0x5C)

    def _heading_run(para, text, size, color=None, bold=True):
        run = para.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color

    # ── Cover Page / Header Branding ────────────────────────
    logo_path = os.path.join(current_app.root_path, "static", "ambifologo.png")
    if os.path.exists(logo_path):
        try:
            logo_para = doc.add_paragraph()
            logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            logo_para.add_run().add_picture(logo_path, width=Inches(2.2))
        except Exception:
            # Keep DOCX export resilient even if image embedding fails.
            pass

    comp_para = doc.add_paragraph()
    comp_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _heading_run(comp_para, "AMBIFO TECHNOLOGY PVT LTD", 12, BRAND_BLUE)

    caps_para = doc.add_paragraph()
    caps_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caps_para.add_run("Cloud Strategy | Migration & Modernization | DevOps | AI/ML").italic = True

    doc.add_paragraph()

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _heading_run(title_para, (document_title_label or "Statement of Work (SOW)").upper(), 22, BRAND_BLUE)

    doc.add_paragraph()

    info_lines = [
        ("Proposal For:", f"{customer.customer_name} — {customer.account_name or ''}"),
        ("Submitted By:", "Ambifo Technology Pvt Ltd"),
        ("Title:", sow_title or f"Cloud Migration SOW — {customer.customer_name}"),
        ("Doc Ref #:", _format_doc_ref_no(doc_ref_no)),
        ("Date:", sow_date or ""),
        ("Version:", sow_version or "1.0"),
        ("Customer Email:", customer.email or ""),
        ("Customer Phone:", customer.phone or "—"),
        ("Cloud Platform:", customer.cloud or "AWS / Azure"),
        ("Website:", "www.ambifo.com"),
        ("Support Email:", "support@ambifo.com"),
        ("Support Phone:", "+91 9148419502"),
        ("Address:", "Building No 674, 18th Main, 3rd Phase, Front of New Land ISRO Quarter, Domlur, Bangalore - 560071"),
        ("LinkedIn:", "https://www.linkedin.com/company/ambifo-technology"),
    ]
    for label, value in info_lines:
        p = doc.add_paragraph()
        run_l = p.add_run(f"{label} ")
        run_l.bold = True
        p.add_run(value)

    doc.add_page_break()

    # ── Body content (already sequence-assembled from template sections) ──
    _html_to_docx_body(doc, content_html)

    # ── Confidentiality footer ───────────────────────────────
    doc.add_paragraph()
    footer_p = doc.add_paragraph()
    footer_run = footer_p.add_run(
        "CONFIDENTIAL — This document and any attachments are confidential and intended solely "
        "for the use of the individual or entity to which it is addressed. "
        "Ambifo Technology Pvt Ltd | support@ambifo.com | www.ambifo.com | +91 9148419502"
    )
    footer_run.italic = True
    footer_run.font.size = Pt(8)

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


def _generate_sow_pdf(customer, sow_title, sow_date, sow_version, content_html, document_title_label="Statement of Work (SOW)", doc_ref_no=""):
    """Build and return a PDF file as bytes."""
    from io import BytesIO
    from bs4 import BeautifulSoup
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas

    def _normalize_text(value):
        return re.sub(r"\s+", " ", (value or "").strip())

    def _html_to_lines(raw_html):
        soup = BeautifulSoup(raw_html or "", "html.parser")
        lines = []
        for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "blockquote"]):
            txt = _normalize_text(el.get_text(" ", strip=True))
            if not txt:
                continue
            if el.name == "li":
                txt = f"- {txt}"
            lines.append(txt)
        if not lines:
            fallback = _normalize_text(soup.get_text(" ", strip=True))
            if fallback:
                lines.append(fallback)
        return lines

    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    page_w, page_h = A4
    left = 46
    right = 46
    top = page_h - 48
    bottom = 52
    y = top
    max_w = page_w - left - right

    def _new_page():
        nonlocal y
        pdf.showPage()
        y = top

    def _draw_wrapped(text, font_name="Helvetica", font_size=11, leading=15):
        nonlocal y
        chunks = simpleSplit(text, font_name, font_size, max_w) or [""]
        for chunk in chunks:
            if y < bottom:
                _new_page()
            pdf.setFont(font_name, font_size)
            pdf.drawString(left, y, chunk)
            y -= leading

    pdf.setTitle(_normalize_text(sow_title) or "SOW")

    _draw_wrapped("AMBIFO TECHNOLOGY PVT LTD", "Helvetica-Bold", 12, 16)
    _draw_wrapped("Cloud Strategy | Migration & Modernization | DevOps | AI/ML", "Helvetica-Oblique", 10, 14)
    y -= 6
    _draw_wrapped((document_title_label or "Statement of Work (SOW)").upper(), "Helvetica-Bold", 16, 20)
    y -= 4

    info_lines = [
        f"Proposal For: {_normalize_text(customer.customer_name)} - {_normalize_text(customer.account_name)}",
        "Submitted By: Ambifo Technology Pvt Ltd",
        f"Title: {_normalize_text(sow_title) or f'Cloud Migration SOW - {_normalize_text(customer.customer_name)}'}",
        f"Doc Ref #: {_format_doc_ref_no(doc_ref_no)}",
        f"Date: {_normalize_text(sow_date)}",
        f"Version: {_normalize_text(sow_version) or '1.0'}",
        f"Customer Email: {_normalize_text(customer.email)}",
        f"Customer Phone: {_normalize_text(customer.phone) or '-'}",
        f"Cloud Platform: {_normalize_text(customer.cloud) or 'AWS / Azure'}",
        "Website: www.ambifo.com",
        "Support Email: support@ambifo.com",
        "Support Phone: +91 9148419502",
    ]
    for ln in info_lines:
        _draw_wrapped(ln, "Helvetica", 10, 14)

    y -= 10
    _draw_wrapped("Document Content", "Helvetica-Bold", 12, 16)
    y -= 2
    for line in _html_to_lines(content_html):
        _draw_wrapped(line, "Helvetica", 10, 14)

    y -= 8
    _draw_wrapped(
        "CONFIDENTIAL - This document and any attachments are confidential and intended solely for the use of the individual or entity to which it is addressed.",
        "Helvetica-Oblique",
        8,
        11,
    )

    pdf.save()
    output.seek(0)
    return output.getvalue()


def _generate_html_like_pdf(title, content_html, header_context=None, footer_context=None):
    """Generate a PDF that preserves editor HTML layout as closely as possible."""
    try:
        from playwright.sync_api import sync_playwright  # noqa: PLC0415
    except ImportError as exc:
        raise RuntimeError("playwright is not installed. Run: pip install playwright && playwright install chromium") from exc

    safe_title = html.escape((title or "Document").strip() or "Document")
    header_context = header_context or {}
    footer_context = footer_context or {}
    raw_html = (content_html or "").strip()
    is_full_document = bool(re.search(r"<html[\s>]", raw_html, flags=re.IGNORECASE))

    app_base = (current_app.config.get("APP_BASE_URL", "http://127.0.0.1:5000") or "http://127.0.0.1:5000").rstrip("/") + "/"

    def _inject_base_tag(doc_html):
        if re.search(r"<base\s+href=", doc_html, flags=re.IGNORECASE):
            return doc_html
        return re.sub(
            r"<head(\s[^>]*)?>",
            lambda m: f"{m.group(0)}\n<base href=\"{html.escape(app_base)}\">",
            doc_html,
            count=1,
            flags=re.IGNORECASE,
        )

    def _inject_pdf_footer_fix(doc_html):
        # PDF-only guardrails for legacy saved documents where footer text may clip on even pages.
        fix_css = """
<style id=\"ambifo-pdf-footer-fix\">
.a4-page {
  display: grid !important;
  grid-template-rows: var(--page-header-space, 122px) minmax(0, 1fr) auto !important;
  width: 794px !important;
  height: 1123px !important;
  min-height: 1123px !important;
  max-height: 1123px !important;
  box-sizing: border-box !important;
  overflow: hidden !important;
  break-after: page !important;
  page-break-after: always !important;
}
.a4-page:last-child {
  break-after: auto !important;
  page-break-after: auto !important;
}
.pg-body {
  min-height: 0 !important;
  overflow: hidden !important;
  box-sizing: border-box !important;
}
.pg-footer {
  overflow: visible !important;
  box-sizing: border-box !important;
  align-items: flex-start !important;
  gap: 8px !important;
  padding-top: 6px !important;
  padding-bottom: 6px !important;
  font-size: 9px !important;
  line-height: 1.25 !important;
}
.pg-footer .left,
.pg-footer .mid,
.pg-footer .right {
  line-height: 1.2 !important;
  overflow-wrap: anywhere !important;
  word-break: break-word !important;
}
</style>
""".strip()
        if re.search(r"id=[\"']ambifo-pdf-footer-fix[\"']", doc_html, flags=re.IGNORECASE):
            return doc_html
        if re.search(r"</head>", doc_html, flags=re.IGNORECASE):
            return re.sub(r"</head>", f"{fix_css}\n</head>", doc_html, count=1, flags=re.IGNORECASE)
        return f"{fix_css}\n{doc_html}"

    if is_full_document:
        full_html = _inject_pdf_footer_fix(_inject_base_tag(raw_html))
    else:
        body_html = raw_html
        full_html = f"""
<!DOCTYPE html>
<html>
<head>
    <meta charset=\"utf-8\">
    <base href=\"{html.escape(app_base)}\">
    <title>{safe_title}</title>
    <style>
        @page {{ size: A4; margin: 18mm 14mm; }}
        html, body {{ margin: 0; padding: 0; }}
        body {{ font-family: 'Segoe UI', Arial, sans-serif; font-size: 10.5pt; line-height: 1.7; color: #1a202c; }}
        h1 {{ font-size: 1.9em; color: #0f4c3a; margin: 1em 0 .35em; }}
        h2 {{ font-size: 1.45em; color: #0f4c3a; margin: .85em 0 .3em; }}
        h3 {{ font-size: 1.15em; color: #0f4c3a; margin: .8em 0 .28em; }}
        h4 {{ font-size: 1em; color: #0f4c3a; margin: .75em 0 .25em; }}
        p {{ margin: 0 0 .55em; }}
        ul, ol {{ padding-left: 1.5em; margin-bottom: .55em; }}
        li {{ margin-bottom: .2em; }}
        blockquote {{ border-left: 3px solid #94a3b8; padding-left: 14px; color: #475569; margin: .6em 0; }}
        code {{ background: #f1f5f9; border-radius: 4px; padding: 1px 5px; font-size: .88em; }}
        pre {{ background: #1e293b; color: #e2e8f0; border-radius: 6px; padding: 12px 16px; margin-bottom: .7em; overflow: auto; }}
        pre code {{ background: none; padding: 0; }}
        table {{ width: 100%; border-collapse: collapse; margin-bottom: .8em; }}
        td, th {{ border: 1px solid #cbd5e1; padding: 6px 10px; font-size: .9em; vertical-align: top; }}
        th {{ background: #f1f5f9; font-weight: 700; }}
        hr {{ border: none; border-top: 2px solid #e2e8f0; margin: 1em 0; }}
        img {{ max-width: 100%; height: auto; border-radius: 4px; }}
        a {{ color: #0d7f5f; text-decoration: underline; }}
        mark {{ background: #fef08a; border-radius: 2px; padding: 0 2px; }}
    </style>
</head>
<body>
{body_html}
</body>
</html>
"""

    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            page = browser.new_page()
            page.set_content(full_html, wait_until="networkidle")
            page.emulate_media(media="screen" if is_full_document else "print")
            page.evaluate(
                """
                () => Promise.all([
                  document.fonts ? document.fonts.ready : Promise.resolve(),
                  Promise.all(Array.from(document.images || []).map(img => {
                    if (img.complete) return Promise.resolve();
                    return new Promise(resolve => {
                      img.addEventListener('load', resolve, { once: true });
                      img.addEventListener('error', resolve, { once: true });
                    });
                  })),
                ])
                """
            )
            pdf_margin = {"top": "0mm", "right": "0mm", "bottom": "0mm", "left": "0mm"} if is_full_document else {"top": "18mm", "right": "14mm", "bottom": "18mm", "left": "14mm"}
            pdf_scale = 0.98 if is_full_document else 1
            pdf_bytes = page.pdf(
                format="A4",
                print_background=True,
                prefer_css_page_size=True,
                display_header_footer=False,
                margin=pdf_margin,
                scale=pdf_scale,
            )
            browser.close()
            return pdf_bytes
    except Exception as exc:
        raise RuntimeError(f"Chromium PDF rendering failed: {exc}. Run: playwright install chromium") from exc


def _sanitize_sow_export_html(content_html):
    """Remove editor-only UI artifacts from HTML before DOCX/email export."""
    text = (content_html or "").strip()
    if not text:
        return ""

    # Remove known editor UI fragments if they were accidentally persisted in HTML content.
    cleanup_patterns = [
        r"<[^>]*id=[\"']selectedDiagramIds[\"'][^>]*>[\s\S]*?</[^>]+>",
        r"<[^>]*id=[\"']macroWarning[\"'][^>]*>[\s\S]*?</[^>]+>",
        r"<[^>]*id=[\"']aiDocObjective[\"'][^>]*>",
        r"<[^>]*class=[\"'][^\"']*editor-label[^\"']*[\"'][^>]*>[\s\S]*?</[^>]+>",
        r"<label[^>]*for=[\"']selectedDiagramIds[\"'][^>]*>[\s\S]*?</label>",
    ]
    for pattern in cleanup_patterns:
        text = re.sub(pattern, "", text, flags=re.IGNORECASE)

    # Remove any plain-text remnants of the UI helper lines.
    text = re.sub(r"Document\s*Content\s*✏️\s*Editable", "", text, flags=re.IGNORECASE)
    text = re.sub(
        r"AI\s*objective\s*\(optional\)\s*:\s*e\.g\.\s*Build migration SOW with phases, risks, and landing-zone architecture",
        "",
        text,
        flags=re.IGNORECASE,
    )
    return text.strip()


def _extract_template_macros(template):
    """Return sorted list of macro keys used in template subject + body (excluding auto-computed ones)."""
    auto_keys = {"today", "gathering_form_link", "gathering_access_key", "gathering_expires_at",
                 "meeting_availability_form_link", "meeting_availability_expires_at", "selected_diagrams_html"}
    combined = f"{template.subject_template or ''} {template.body_template or ''}"
    found = re.findall(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}", combined)
    return sorted({k for k in found if k not in auto_keys})


def _lead_sample_columns():
    return ["lead_name", "email", "phone", "company", "city", "source", "tags", "notes"]


def _upsert_lead_from_row(row, uploaded_from=None):
    email = (row.get("email") or "").strip().lower()
    if not email or not _is_valid_email(email):
        return False

    lead = Lead.query.filter_by(email=email).first()
    if not lead:
        lead = Lead(email=email)
        db.session.add(lead)

    lead.lead_name = (row.get("lead_name") or lead.lead_name or "").strip() or None
    lead.phone = (row.get("phone") or lead.phone or "").strip() or None
    lead.company = (row.get("company") or lead.company or "").strip() or None
    lead.city = (row.get("city") or lead.city or "").strip() or None
    lead.source = (row.get("source") or lead.source or uploaded_from or "").strip() or None
    lead.tags = (row.get("tags") or lead.tags or "").strip() or None
    lead.notes = (row.get("notes") or lead.notes or "").strip() or None
    try:
        lead.raw_payload = json.dumps(row, ensure_ascii=True)
    except Exception:
        lead.raw_payload = None
    return True


def _lead_macro_values(lead):
    return {
        "today": datetime.utcnow().date().isoformat(),
        "lead_name": lead.lead_name or "",
        "customer_name": lead.lead_name or "",
        "account_name": lead.company or "",
        "company": lead.company or "",
        "email": lead.email or "",
        "phone": lead.phone or "",
        "city": lead.city or "",
        "source": lead.source or "",
        "tags": lead.tags or "",
    }


@crm_bp.route("/configuration/sow-template-sections/<int:template_id>")
def configuration_sow_template_sections(template_id):
    # Backward-compatible alias used by older templates; section management now lives on /configuration.
    section_id = request.args.get("section_id", type=int)
    target = url_for("crm.configuration", sow_template_id=template_id)
    if section_id:
        target = f"{target}&section_id={section_id}"
    return redirect(f"{target}#sow-master-config")


@crm_bp.route("/configuration", methods=["GET", "POST"])
def configuration():
    _ensure_default_statuses()
    _ensure_default_segments()
    _ensure_default_cloud_operators()
    _ensure_default_sow_master_template()

    if request.method == "POST":
        action = (request.form.get("action") or "").strip()

        admin_action_response = handle_configuration_admin_action(action)
        if admin_action_response is not None:
            return admin_action_response

        storage_action_response = handle_configuration_storage_action(
            action,
            persist_settings=_persist_document_storage_settings,
            apply_runtime=_apply_document_storage_runtime,
            log_history=_log_opportunity_history,
        )
        if storage_action_response is not None:
            return storage_action_response

        if action == "save_smtp":
            SystemSetting.set_value("smtp.host", (request.form.get("smtp_host") or "").strip())
            SystemSetting.set_value("smtp.port", (request.form.get("smtp_port") or "587").strip())
            SystemSetting.set_value("smtp.username", (request.form.get("smtp_username") or "").strip())
            SystemSetting.set_value("smtp.password", (request.form.get("smtp_password") or "").strip())
            SystemSetting.set_value("smtp.use_tls", "true" if request.form.get("smtp_use_tls") else "false")
            SystemSetting.set_value("smtp.mail_from", (request.form.get("smtp_mail_from") or "").strip())
            db.session.commit()
            flash("SMTP configuration saved.", "success")

        elif action == "test_smtp":
            test_to = (request.form.get("test_email_to") or "").strip()
            if not test_to or "@" not in test_to:
                flash("Please enter a valid test recipient email.", "error")
            else:
                email_service = EmailService(current_app)
                tpl, subject, body = _render_system_email_template(
                    "smtp_test",
                    {
                        "today": datetime.utcnow().date().isoformat(),
                        "test_email_to": test_to,
                    },
                )
                result = email_service.send_html_email(test_to, subject, body)
                db.session.add(EmailLog(
                    customer_id=None,
                    template_id=tpl.id if tpl else None,
                    recipient_email=test_to,
                    email_type="smtp-test",
                    subject=subject,
                    body=body,
                    status=result.status,
                    error_message=result.error,
                ))
                try:
                    db.session.commit()
                except Exception:
                    db.session.rollback()
                    if result.success or result.status == "dev-mode":
                        flash(
                            "SMTP test email sent, but log insert failed due to old DB schema. Restart app once to auto-fix.",
                            "error",
                        )
                    else:
                        flash(f"SMTP test failed: {result.error}", "error")
                    return redirect(url_for("crm.configuration") + "#smtp-config")
                if result.success or result.status == "dev-mode":
                    flash(f"Test email sent to {test_to} (status: {result.status}).", "success")
                else:
                    flash(f"SMTP test failed: {result.error}", "error")

        elif action == "add_template":
            name = (request.form.get("name") or "").strip()
            subject_template = (request.form.get("subject_template") or "").strip()
            body_template = _sanitize_email_template_body(request.form.get("body_template") or "")
            if not name or not subject_template or not body_template:
                flash("Template name, subject, and body are required.", "error")
            elif EmailTemplate.query.filter_by(name=name).first():
                flash("Template name already exists.", "error")
            else:
                db.session.add(EmailTemplate(
                    name=name,
                    subject_template=subject_template,
                    body_template=body_template,
                ))
                db.session.commit()
                flash("Email template created.", "success")
            return redirect(url_for("crm.configuration") + "#email-template-config")

        elif action == "edit_template":
            tpl_id = request.form.get("template_id", type=int)
            tpl = EmailTemplate.query.get(tpl_id)
            if not tpl:
                flash("Template not found.", "error")
                return redirect(url_for("crm.configuration") + "#email-template-config")

            name = (request.form.get("name") or "").strip()
            subject_template = (request.form.get("subject_template") or "").strip()
            body_template = _sanitize_email_template_body(request.form.get("body_template") or "")
            if not name or not subject_template or not body_template:
                flash("Template name, subject, and body are required.", "error")
                return redirect(url_for("crm.configuration") + "#email-template-config")

            duplicate = EmailTemplate.query.filter(
                EmailTemplate.name == name,
                EmailTemplate.id != tpl.id,
            ).first()
            if duplicate:
                flash("Another template already uses that name.", "error")
                return redirect(url_for("crm.configuration") + "#email-template-config")

            tpl.name = name
            tpl.subject_template = subject_template
            tpl.body_template = body_template
            db.session.commit()
            flash("Email template updated.", "success")
            return redirect(url_for("crm.configuration") + "#email-template-config")

        elif action == "delete_template":
            tpl_id = request.form.get("template_id", type=int)
            tpl = EmailTemplate.query.get(tpl_id)
            if not tpl:
                flash("Template not found.", "error")
            elif _is_system_template_name(tpl.name):
                flash("System templates cannot be deleted. You can edit them.", "error")
            else:
                db.session.delete(tpl)
                db.session.commit()
                flash("Email template deleted.", "success")
            return redirect(url_for("crm.configuration") + "#email-template-config")

        elif action == "test_email_template":
            template_id = request.form.get("template_id", type=int)
            test_to = (request.form.get("test_email_to") or "").strip()
            customer_id = request.form.get("test_customer_id", type=int)

            if not template_id:
                flash("Please select an email template.", "error")
                return redirect(url_for("crm.configuration") + "#email-template-config")
            if not test_to or "@" not in test_to:
                flash("Please enter a valid recipient email for test.", "error")
                return redirect(url_for("crm.configuration") + "#email-template-config")

            template = EmailTemplate.query.get(template_id)
            if not template:
                flash("Template not found.", "error")
                return redirect(url_for("crm.configuration") + "#email-template-config")

            customer = Customer.query.get(customer_id) if customer_id else None
            template_text = f"{template.subject_template or ''}\n{template.body_template or ''}"
            needs_customer_context = (
                "{{meeting_availability_form_link}}" in template_text
                or "{{meeting_availability_expires_at}}" in template_text
            )

            if needs_customer_context and not customer:
                flash(
                    "This template needs opportunity/customer context. Please choose a customer for test send.",
                    "error",
                )
                return redirect(url_for("crm.configuration") + "#email-template-config")

            if customer:
                macro_values = _build_template_macro_values_for_customer(customer, template, recipient_email=test_to)
                rendered_subject = render_macros(template.subject_template, macro_values)
                rendered_body = render_macros(template.body_template, macro_values)
            else:
                rendered_subject = template.subject_template
                rendered_body = template.body_template

            email_service = EmailService(current_app)
            result = email_service.send_html_email(test_to, rendered_subject, rendered_body)
            db.session.add(EmailLog(
                customer_id=customer.id if customer else None,
                template_id=template.id,
                recipient_email=test_to,
                email_type="template-test",
                subject=rendered_subject,
                body=rendered_body,
                status=result.status,
                error_message=result.error,
            ))
            db.session.commit()

            if result.success or result.status == "dev-mode":
                flash(
                    f"Template test email sent to {test_to} (status: {result.status}).",
                    "success",
                )
            else:
                flash(f"Template test failed: {result.error}", "error")
            return redirect(url_for("crm.configuration") + "#email-template-config")

        elif action == "save_teams":
            SystemSetting.set_value("teams.tenant_id", (request.form.get("teams_tenant_id") or "").strip())
            SystemSetting.set_value("teams.client_id", (request.form.get("teams_client_id") or "").strip())
            SystemSetting.set_value("teams.client_secret", (request.form.get("teams_client_secret") or "").strip())
            SystemSetting.set_value("teams.organizer_id", (request.form.get("teams_organizer_id") or "").strip())
            SystemSetting.set_value(
                "teams.default_duration_minutes",
                (request.form.get("teams_default_duration_minutes") or "60").strip(),
            )
            db.session.commit()
            flash("Teams configuration saved.", "success")

        elif action == "save_meeting_settings":
            raw_days = (request.form.get("meeting_availability_link_expiry_days") or "7").strip()
            try:
                expiry_days = int(raw_days)
            except ValueError:
                flash("Availability link expiry must be a number.", "error")
                return redirect(url_for("crm.configuration"))

            if expiry_days < 1 or expiry_days > 90:
                flash("Availability link expiry must be between 1 and 90 days.", "error")
                return redirect(url_for("crm.configuration"))

            SystemSetting.set_value("meeting.availability_link_expiry_days", str(expiry_days))
            db.session.commit()
            flash("Meeting availability settings saved.", "success")

        elif action == "save_app_settings":
            app_base_url = (request.form.get("app_base_url") or "").strip()
            if not app_base_url.startswith("http://") and not app_base_url.startswith("https://"):
                flash("App Base URL must start with http:// or https://", "error")
                return redirect(url_for("crm.configuration"))

            SystemSetting.set_value("app.base_url", app_base_url.rstrip("/"))
            db.session.commit()
            flash("Application URL settings saved.", "success")

        elif action == "save_sow_brand_config":
            def _parse_px(field_name, default_value):
                raw = (request.form.get(field_name) or str(default_value)).strip()
                try:
                    value = int(raw)
                except (TypeError, ValueError):
                    value = int(default_value)
                return max(0, min(value, 300))

            SystemSetting.set_value("sow.brand.company_name",   (request.form.get("sow_company_name") or "").strip())
            SystemSetting.set_value("sow.brand.tagline_chips",  (request.form.get("sow_tagline_chips") or "").strip())
            SystemSetting.set_value("sow.brand.website",        (request.form.get("sow_website") or "").strip())
            SystemSetting.set_value("sow.brand.email",          (request.form.get("sow_email") or "").strip())
            SystemSetting.set_value("sow.brand.phone",          (request.form.get("sow_phone") or "").strip())
            SystemSetting.set_value("sow.brand.address",        (request.form.get("sow_address") or "").strip())
            SystemSetting.set_value("sow.brand.confidential",   (request.form.get("sow_confidential") or "").strip())
            SystemSetting.set_value("sow.brand.header_top_px",  str(_parse_px("sow_header_top_px", 122)))
            SystemSetting.set_value("sow.brand.footer_bottom_px", str(_parse_px("sow_footer_bottom_px", 80)))
            db.session.commit()
            flash("SOW header/footer brand settings saved.", "success")
            return redirect(url_for("crm.configuration") + "#sow-master-config")

        elif action == "save_ai_settings":
            active_provider = (request.form.get("ai_active_provider") or "openai").strip().lower()
            current_active_provider = (
                (SystemSetting.get_value("ai.active_provider", None)
                or SystemSetting.get_value("ai.provider", "openai")
                or "openai")
                .strip()
                .lower()
            )
            # If provider selection changed, force activation to keep one active provider deterministic.
            activate_selected_provider = bool(request.form.get("ai_activate_selected_provider")) or (
                active_provider != current_active_provider
            )
            current_openai_api_key = (SystemSetting.get_value("ai.openai.api_key", "") or "").strip()
            current_openai_model = (SystemSetting.get_value("ai.openai.model", "gpt-4o-mini") or "gpt-4o-mini").strip()
            current_openai_base_url = (SystemSetting.get_value("ai.openai.base_url", "https://api.openai.com/v1") or "https://api.openai.com/v1").strip()
            current_gemini_api_key = (SystemSetting.get_value("ai.gemini.api_key", "") or "").strip()
            current_gemini_model = (SystemSetting.get_value("ai.gemini.model", "gemini-1.5-flash") or "gemini-1.5-flash").strip()
            current_gemini_base_url = (SystemSetting.get_value("ai.gemini.base_url", "https://generativelanguage.googleapis.com/v1beta") or "https://generativelanguage.googleapis.com/v1beta").strip()

            openai_api_key = (request.form.get("openai_api_key") or current_openai_api_key).strip()
            openai_model = (request.form.get("openai_model") or current_openai_model or "gpt-4o-mini").strip()
            openai_base_url = (request.form.get("openai_base_url") or current_openai_base_url or "https://api.openai.com/v1").strip()
            gemini_api_key = (request.form.get("gemini_api_key") or current_gemini_api_key).strip()
            gemini_model = (request.form.get("gemini_model") or current_gemini_model or "gemini-1.5-flash").strip()
            gemini_base_url = (request.form.get("gemini_base_url") or current_gemini_base_url or "https://generativelanguage.googleapis.com/v1beta").strip()

            if active_provider not in {"openai", "gemini"}:
                flash("Active provider must be either openai or gemini.", "error")
                return redirect(url_for("crm.configuration"))
            if active_provider == "openai":
                if not openai_base_url.startswith("http://") and not openai_base_url.startswith("https://"):
                    flash("OpenAI Base URL must start with http:// or https://", "error")
                    return redirect(url_for("crm.configuration") + "#ai-config")
                if not openai_model:
                    flash("OpenAI model is required.", "error")
                    return redirect(url_for("crm.configuration") + "#ai-config")
            else:
                if not gemini_base_url.startswith("http://") and not gemini_base_url.startswith("https://"):
                    flash("Gemini Base URL must start with http:// or https://", "error")
                    return redirect(url_for("crm.configuration") + "#ai-config")
                if not gemini_model:
                    flash("Gemini model is required.", "error")
                    return redirect(url_for("crm.configuration") + "#ai-config")

            activation_warning = None
            if activate_selected_provider and active_provider == "openai" and not openai_api_key:
                activation_warning = "OpenAI is activated, but API key is empty. AI generation will fail until you add it."
            if activate_selected_provider and active_provider == "gemini" and not gemini_api_key:
                activation_warning = "Gemini is activated, but API key is empty. AI generation will fail until you add it."

            SystemSetting.set_value("ai.openai.api_key", openai_api_key)
            SystemSetting.set_value("ai.openai.model", openai_model)
            SystemSetting.set_value("ai.openai.base_url", openai_base_url.rstrip("/"))
            SystemSetting.set_value("ai.gemini.api_key", gemini_api_key)
            SystemSetting.set_value("ai.gemini.model", gemini_model)
            SystemSetting.set_value("ai.gemini.base_url", gemini_base_url.rstrip("/"))

            if activate_selected_provider:
                SystemSetting.set_value("ai.active_provider", active_provider)
                SystemSetting.set_value("ai.provider", active_provider)

                # Legacy compatibility key for old code paths.
                if active_provider == "openai":
                    SystemSetting.set_value("ai.api_key", openai_api_key)
                else:
                    SystemSetting.set_value("ai.api_key", gemini_api_key)

            db.session.commit()
            if activate_selected_provider:
                flash(
                    f"AI settings saved. Activated provider: {active_provider}. Previous provider is now inactive.",
                    "success",
                )
                if activation_warning:
                    flash(activation_warning, "error")
            else:
                flash("AI settings saved. Active provider unchanged.", "success")

        elif action == "activate_ai_provider":
            provider = (request.form.get("provider") or "").strip().lower()
            if provider not in {"openai", "gemini"}:
                flash("Provider must be either openai or gemini.", "error")
                return redirect(url_for("crm.configuration") + "#ai-config")

            openai_api_key = (SystemSetting.get_value("ai.openai.api_key", "") or "").strip()
            gemini_api_key = (SystemSetting.get_value("ai.gemini.api_key", "") or "").strip()

            activation_warning = None
            if provider == "openai" and not openai_api_key:
                activation_warning = "OpenAI is activated, but API key is empty. AI generation will fail until you add it."
            if provider == "gemini" and not gemini_api_key:
                activation_warning = "Gemini is activated, but API key is empty. AI generation will fail until you add it."

            SystemSetting.set_value("ai.active_provider", provider)
            SystemSetting.set_value("ai.provider", provider)

            # Legacy compatibility key for old code paths.
            if provider == "openai":
                SystemSetting.set_value("ai.api_key", openai_api_key)
            else:
                SystemSetting.set_value("ai.api_key", gemini_api_key)

            db.session.commit()
            flash(f"Active AI provider switched to {provider}.", "success")
            if activation_warning:
                flash(activation_warning, "error")
            return redirect(url_for("crm.configuration") + "#ai-config")

        elif action == "add_sow_master_template":
            template_name = (request.form.get("sow_template_name") or "").strip()
            if not template_name:
                flash("Template name is required.", "error")
                return redirect(url_for("crm.configuration") + "#sow-master-config")

            key = _generate_unique_sow_template_key(template_name)

            master = _create_master_template(template_name, key)
            db.session.add(master)
            db.session.commit()
            _ensure_template_sections(master)
            flash("Document template added.", "success")
            return redirect(
                url_for("crm.configuration", sow_template_id=master.id)
                + "#sow-master-config"
            )

        elif action == "duplicate_sow_master_template":
            source_id = request.form.get("sow_template_id", type=int)
            source = SOWMasterTemplate.query.get(source_id)
            if not source:
                flash("Selected document template was not found.", "error")
                return redirect(url_for("crm.configuration") + "#sow-master-config")

            new_name = f"{source.template_name} Copy"
            new_key = _generate_unique_sow_template_key(new_name)
            duplicate = _create_master_template(new_name, new_key)
            duplicate.content_html = source.content_html
            duplicate.updated_by = _actor_name()
            duplicate.updated_at = datetime.utcnow()

            db.session.add(duplicate)
            db.session.commit()
            source_sections = _ensure_template_sections(source)
            for section in source_sections:
                db.session.add(
                    SOWMasterTemplateSection(
                        template_id=duplicate.id,
                        section_name=section.section_name,
                        sequence_no=section.sequence_no,
                        content_html=section.content_html,
                    )
                )
            db.session.commit()
            flash("Document template duplicated.", "success")
            return redirect(
                url_for("crm.configuration", sow_template_id=duplicate.id)
                + "#sow-master-config"
            )

        elif action == "delete_sow_master_template":
            template_id = request.form.get("sow_template_id", type=int)
            template_obj = SOWMasterTemplate.query.get(template_id)
            if not template_obj:
                flash("Selected document template was not found.", "error")
                return redirect(url_for("crm.configuration") + "#sow-master-config")

            total_templates = SOWMasterTemplate.query.count()
            if total_templates <= 1:
                flash("At least one document template must exist.", "error")
                return redirect(
                    url_for("crm.configuration", sow_template_id=template_obj.id)
                    + "#sow-master-config"
                )

            usage_count = CustomerSOW.query.filter_by(master_template_id=template_obj.id).count()
            if usage_count > 0:
                flash("Cannot delete template because it is already used by opportunities.", "error")
                return redirect(
                    url_for("crm.configuration", sow_template_id=template_obj.id)
                    + "#sow-master-config"
                )

            fallback = SOWMasterTemplate.query.filter(SOWMasterTemplate.id != template_obj.id).order_by(SOWMasterTemplate.id.asc()).first()
            SOWMasterTemplateSection.query.filter_by(template_id=template_obj.id).delete(synchronize_session=False)
            db.session.delete(template_obj)
            db.session.commit()
            flash("Document template deleted.", "success")
            return redirect(
                url_for("crm.configuration", sow_template_id=fallback.id if fallback else None)
                + "#sow-master-config"
            )

        elif action == "add_status":
            name = (request.form.get("status_name") or "").strip()
            color = (request.form.get("status_color") or "#8a8f98").strip()

            if not name:
                flash("Status name is required.", "error")
            elif OpportunityStatus.query.filter_by(name=name).first():
                flash("Status already exists.", "error")
            else:
                db.session.add(OpportunityStatus(name=name, color=color, is_active=True))
                db.session.commit()
                flash("Status added.", "success")

        elif action == "update_status":
            status_id = request.form.get("status_id", type=int)
            name = (request.form.get("status_name") or "").strip()
            color = (request.form.get("status_color") or "#8a8f98").strip()

            status_obj = OpportunityStatus.query.get(status_id)
            if not status_obj:
                flash("Status not found.", "error")
            elif not name:
                flash("Status name is required.", "error")
            else:
                duplicate = OpportunityStatus.query.filter(
                    OpportunityStatus.name == name,
                    OpportunityStatus.id != status_obj.id,
                ).first()
                if duplicate:
                    flash("Another status already uses this name.", "error")
                else:
                    old_name = status_obj.name
                    status_obj.name = name
                    status_obj.color = color
                    if old_name != name:
                        Customer.query.filter(Customer.deal_status == old_name).update(
                            {Customer.deal_status: name},
                            synchronize_session=False,
                        )
                    db.session.commit()
                    flash("Status updated.", "success")

        elif action == "delete_status":
            status_id = request.form.get("status_id", type=int)
            status_obj = OpportunityStatus.query.get(status_id)
            if not status_obj:
                flash("Status not found.", "error")
            else:
                Customer.query.filter(Customer.deal_status == status_obj.name).update(
                    {Customer.deal_status: ""},
                    synchronize_session=False,
                )
                db.session.delete(status_obj)
                db.session.commit()
                flash("Status deleted.", "success")

        elif action == "add_segment":
            name = (request.form.get("segment_name") or "").strip()
            credit_percentage_customer_raw = (request.form.get("segment_credit_percentage_customer") or "").strip()
            credit_percentage_ambifo_raw = (request.form.get("segment_credit_percentage_ambifo") or "").strip()
            credit_basis_mrr = bool(request.form.get("segment_credit_basis_mrr"))
            credit_basis_arr = bool(request.form.get("segment_credit_basis_arr"))

            if credit_basis_mrr and credit_basis_arr:
                flash("Select only one basis: MRR or ARR.", "error")
                return redirect(url_for("crm.configuration"))
            if not credit_basis_mrr and not credit_basis_arr:
                credit_basis_mrr = True

            credit_percentage_customer = None
            if credit_percentage_customer_raw:
                try:
                    credit_percentage_customer = float(credit_percentage_customer_raw)
                except ValueError:
                    flash("Credit % To Customer must be a valid number.", "error")
                    return redirect(url_for("crm.configuration"))
                if credit_percentage_customer < 0:
                    flash("Credit % To Customer cannot be negative.", "error")
                    return redirect(url_for("crm.configuration"))

            credit_percentage_ambifo = None
            if credit_percentage_ambifo_raw:
                try:
                    credit_percentage_ambifo = float(credit_percentage_ambifo_raw)
                except ValueError:
                    flash("Credit % To Ambifo must be a valid number.", "error")
                    return redirect(url_for("crm.configuration"))
                if credit_percentage_ambifo < 0:
                    flash("Credit % To Ambifo cannot be negative.", "error")
                    return redirect(url_for("crm.configuration"))

            if not name:
                flash("Segment name is required.", "error")
            elif OpportunitySegment.query.filter_by(name=name).first():
                flash("Segment already exists.", "error")
            else:
                db.session.add(
                    OpportunitySegment(
                        name=name,
                        credit_percentage_customer=credit_percentage_customer,
                        credit_percentage_ambifo=credit_percentage_ambifo,
                        credit_basis_mrr=credit_basis_mrr,
                        credit_basis_arr=credit_basis_arr,
                        is_active=True,
                    )
                )
                db.session.commit()
                flash("Segment added.", "success")

        elif action == "update_segment":
            segment_id = request.form.get("segment_id", type=int)
            name = (request.form.get("segment_name") or "").strip()
            credit_percentage_customer_raw = (request.form.get("segment_credit_percentage_customer") or "").strip()
            credit_percentage_ambifo_raw = (request.form.get("segment_credit_percentage_ambifo") or "").strip()
            credit_basis_mrr = bool(request.form.get("segment_credit_basis_mrr"))
            credit_basis_arr = bool(request.form.get("segment_credit_basis_arr"))

            if credit_basis_mrr and credit_basis_arr:
                flash("Select only one basis: MRR or ARR.", "error")
                return redirect(url_for("crm.configuration"))
            if not credit_basis_mrr and not credit_basis_arr:
                credit_basis_mrr = True

            credit_percentage_customer = None
            if credit_percentage_customer_raw:
                try:
                    credit_percentage_customer = float(credit_percentage_customer_raw)
                except ValueError:
                    flash("Credit % To Customer must be a valid number.", "error")
                    return redirect(url_for("crm.configuration"))
                if credit_percentage_customer < 0:
                    flash("Credit % To Customer cannot be negative.", "error")
                    return redirect(url_for("crm.configuration"))

            credit_percentage_ambifo = None
            if credit_percentage_ambifo_raw:
                try:
                    credit_percentage_ambifo = float(credit_percentage_ambifo_raw)
                except ValueError:
                    flash("Credit % To Ambifo must be a valid number.", "error")
                    return redirect(url_for("crm.configuration"))
                if credit_percentage_ambifo < 0:
                    flash("Credit % To Ambifo cannot be negative.", "error")
                    return redirect(url_for("crm.configuration"))

            segment_obj = OpportunitySegment.query.get(segment_id)
            if not segment_obj:
                flash("Segment not found.", "error")
            elif not name:
                flash("Segment name is required.", "error")
            else:
                duplicate = OpportunitySegment.query.filter(
                    OpportunitySegment.name == name,
                    OpportunitySegment.id != segment_obj.id,
                ).first()
                if duplicate:
                    flash("Another segment already uses this name.", "error")
                else:
                    old_name = segment_obj.name
                    segment_obj.name = name
                    segment_obj.credit_percentage_customer = credit_percentage_customer
                    segment_obj.credit_percentage_ambifo = credit_percentage_ambifo
                    segment_obj.credit_basis_mrr = credit_basis_mrr
                    segment_obj.credit_basis_arr = credit_basis_arr
                    if old_name != name:
                        Customer.query.filter(Customer.segment == old_name).update(
                            {Customer.segment: name},
                            synchronize_session=False,
                        )
                    _recalculate_financials_for_segment(name)
                    db.session.commit()
                    flash("Segment updated.", "success")

        elif action == "delete_segment":
            segment_id = request.form.get("segment_id", type=int)
            segment_obj = OpportunitySegment.query.get(segment_id)
            if not segment_obj:
                flash("Segment not found.", "error")
            else:
                affected_customers = Customer.query.filter(Customer.segment == segment_obj.name).all()
                Customer.query.filter(Customer.segment == segment_obj.name).update(
                    {Customer.segment: ""},
                    synchronize_session=False,
                )
                for customer in affected_customers:
                    if not customer.financial_profile:
                        continue
                    payload = _build_financial_payload(
                        {
                            "expected_mrr": customer.financial_profile.expected_mrr,
                            "expected_arr": customer.financial_profile.expected_arr,
                            "actual_mrr": customer.financial_profile.actual_mrr,
                            "actual_arr": customer.financial_profile.actual_arr,
                        },
                        "",
                    )
                    _apply_financial_payload(customer.id, payload)
                db.session.delete(segment_obj)
                db.session.commit()
                flash("Segment deleted.", "success")

        elif action == "add_cloud_operator":
            name = (request.form.get("cloud_name") or "").strip()

            if not name:
                flash("Cloud operator name is required.", "error")
            elif OpportunityCloudOperator.query.filter_by(name=name).first():
                flash("Cloud operator already exists.", "error")
            else:
                db.session.add(OpportunityCloudOperator(name=name, is_active=True))
                db.session.commit()
                flash("Cloud operator added.", "success")

        elif action == "update_cloud_operator":
            cloud_id = request.form.get("cloud_id", type=int)
            name = (request.form.get("cloud_name") or "").strip()

            cloud_obj = OpportunityCloudOperator.query.get(cloud_id)
            if not cloud_obj:
                flash("Cloud operator not found.", "error")
            elif not name:
                flash("Cloud operator name is required.", "error")
            else:
                duplicate = OpportunityCloudOperator.query.filter(
                    OpportunityCloudOperator.name == name,
                    OpportunityCloudOperator.id != cloud_obj.id,
                ).first()
                if duplicate:
                    flash("Another cloud operator already uses this name.", "error")
                else:
                    old_name = cloud_obj.name
                    cloud_obj.name = name
                    if old_name != name:
                        Customer.query.filter(Customer.cloud == old_name).update(
                            {Customer.cloud: name},
                            synchronize_session=False,
                        )
                    db.session.commit()
                    flash("Cloud operator updated.", "success")

        elif action == "delete_cloud_operator":
            cloud_id = request.form.get("cloud_id", type=int)
            cloud_obj = OpportunityCloudOperator.query.get(cloud_id)
            if not cloud_obj:
                flash("Cloud operator not found.", "error")
            else:
                Customer.query.filter(Customer.cloud == cloud_obj.name).update(
                    {Customer.cloud: ""},
                    synchronize_session=False,
                )
                db.session.delete(cloud_obj)
                db.session.commit()
                flash("Cloud operator deleted.", "success")

        elif action == "add_update_tag":
            name = (request.form.get("tag_name") or "").strip().lower()
            color = (request.form.get("tag_color") or "#6b7280").strip()

            if not name:
                flash("Tag name is required.", "error")
            elif OpportunityUpdateTag.query.filter_by(name=name).first():
                flash("Tag already exists.", "error")
            else:
                db.session.add(OpportunityUpdateTag(name=name, color=color, is_active=True))
                db.session.commit()
                flash("Update tag added.", "success")

        elif action == "update_update_tag":
            tag_id = request.form.get("tag_id", type=int)
            name = (request.form.get("tag_name") or "").strip().lower()
            color = (request.form.get("tag_color") or "#6b7280").strip()

            tag_obj = OpportunityUpdateTag.query.get(tag_id)
            if not tag_obj:
                flash("Tag not found.", "error")
            elif not name:
                flash("Tag name is required.", "error")
            else:
                duplicate = OpportunityUpdateTag.query.filter(
                    OpportunityUpdateTag.name == name,
                    OpportunityUpdateTag.id != tag_obj.id,
                ).first()
                if duplicate:
                    flash("Another tag already uses this name.", "error")
                else:
                    old_name = tag_obj.name
                    tag_obj.name = name
                    tag_obj.color = color
                    if old_name != name:
                        OpportunityHistory.query.filter(OpportunityHistory.tag_name == old_name).update(
                            {OpportunityHistory.tag_name: name},
                            synchronize_session=False,
                        )
                    db.session.commit()
                    flash("Update tag updated.", "success")

        elif action == "delete_update_tag":
            tag_id = request.form.get("tag_id", type=int)
            tag_obj = OpportunityUpdateTag.query.get(tag_id)
            if not tag_obj:
                flash("Tag not found.", "error")
            elif tag_obj.name in {"email", "important", "update", "meeting"}:
                flash("This default tag cannot be deleted.", "error")
            else:
                db.session.delete(tag_obj)
                db.session.commit()
                flash("Update tag deleted.", "success")

        elif action == "restart_app":
            _request_process_restart()
            flash("Restart requested. The service should come back in a few seconds.", "success")

        return redirect(url_for("crm.configuration"))

    smtp_settings = _get_smtp_settings()
    teams_settings = _get_teams_settings()
    meeting_settings = _get_meeting_settings()
    ai_settings = _get_ai_settings()
    app_settings = {
        "base_url": _get_effective_app_base_url(),
    }
    document_storage_settings = _get_document_storage_settings()
    _ensure_default_sow_master_template()
    sow_master_templates = SOWMasterTemplate.query.order_by(SOWMasterTemplate.created_at.asc()).all()
    selected_sow_template_id = request.args.get("sow_template_id", type=int)

    selected_sow_template = None
    if selected_sow_template_id:
        selected_sow_template = next((t for t in sow_master_templates if t.id == selected_sow_template_id), None)
    if not selected_sow_template and sow_master_templates:
        selected_sow_template = sow_master_templates[0]

    selected_sow_template_usage_count = 0
    if selected_sow_template:
        _ensure_template_sections(selected_sow_template)
        selected_sow_template_usage_count = CustomerSOW.query.filter_by(master_template_id=selected_sow_template.id).count()

    statuses = OpportunityStatus.query.order_by(OpportunityStatus.created_at.asc()).all()
    segments = OpportunitySegment.query.order_by(OpportunitySegment.created_at.asc()).all()
    cloud_operators = OpportunityCloudOperator.query.order_by(OpportunityCloudOperator.created_at.asc()).all()
    update_tags = OpportunityUpdateTag.query.order_by(OpportunityUpdateTag.created_at.asc()).all()
    admin_users = User.query.order_by(User.created_at.asc()).all()
    admin_history_notification_subscriptions = load_admin_history_notification_subscriptions()
    notification_tag_names = get_notification_tag_names()
    email_templates = EmailTemplate.query.order_by(EmailTemplate.created_at.desc()).all()
    customers = Customer.query.order_by(Customer.customer_name.asc()).all()
    active_diagrams = (
        CustomerDiagram.query
        .filter_by(is_active=True)
        .order_by(CustomerDiagram.customer_id.asc(), CustomerDiagram.macro_key.asc())
        .all()
    )
    diagram_macro_map = {}
    for diagram in active_diagrams:
        customer_key = str(diagram.customer_id)
        if customer_key not in diagram_macro_map:
            diagram_macro_map[customer_key] = []
        if (diagram.macro_key or "").strip():
            diagram_macro_map[customer_key].append(
                {
                    "macro_key": diagram.macro_key,
                    "diagram_name": diagram.diagram_name,
                }
            )

    return render_template(
        "configuration.html",
        smtp=smtp_settings,
        teams=teams_settings,
        meeting=meeting_settings,
        ai=ai_settings,
        app_settings=app_settings,
        document_storage=document_storage_settings,
        sow_master_templates=sow_master_templates,
        selected_sow_template=selected_sow_template,
        selected_sow_template_usage_count=selected_sow_template_usage_count,
        sow_template_count=len(sow_master_templates),
        sow_brand_cfg=_get_sow_brand_config(),
        statuses=statuses,
        segments=segments,
        cloud_operators=cloud_operators,
        update_tags=update_tags,
        admin_users=admin_users,
        admin_history_notification_subscriptions=admin_history_notification_subscriptions,
        notification_tag_names=notification_tag_names,
        email_templates=email_templates,
        customers=customers,
        diagram_macro_map=diagram_macro_map,
        email_macro_catalog=_email_template_macro_catalog(),
        system_template_names=_system_template_name_set(),
    )


# -- Gathering data: CSV / XLSX import (Servers + FileNAS + Block Storage) ---

# ── Route Module Registration ────────────────────────────────────────────────

register_auth_routes(crm_bp)

register_dashboard_routes(
    crm_bp,
    derive_tag_from_action=_derive_tag_from_action,
    tag_color_map=_tag_color_map,
)

register_ai_routes(
    crm_bp,
    build_ai_customer_context=_build_ai_customer_context,
    get_template_sections=_get_template_sections,
)

register_email_routes(
    crm_bp,
    build_template_macro_values_for_customer=_build_template_macro_values_for_customer,
    customer_alternate_email_list=_customer_alternate_email_list,
    is_valid_email=_is_valid_email,
    log_opportunity_history=_log_opportunity_history,
    merge_cc_lists=_merge_cc_lists,
)

register_gathering_request_routes(
    crm_bp,
    actor_name=_actor_name,
    generate_access_key=_generate_access_key,
    get_effective_app_base_url=_get_effective_app_base_url,
    import_uploaded_file_for_target=_import_uploaded_file_for_target,
    import_uploaded_file_to_all_tables=_import_uploaded_file_to_all_tables,
    is_gathering_request_expired=_is_gathering_request_expired,
    log_opportunity_history=_log_opportunity_history,
    render_system_email_template=_render_system_email_template,
)

register_gathering_data_routes(
    crm_bp,
    detect_rows_type=_detect_rows_type,
    import_block_storage_rows=_import_block_storage_rows,
    import_file_nas_rows=_import_file_nas_rows,
    import_server_rows=_import_server_rows,
    log_opportunity_history=_log_opportunity_history,
    read_csv_rows=_read_csv_rows,
    read_xlsx_rows_by_sheet=_read_xlsx_rows_by_sheet,
    to_float=_to_float,
    to_int=_to_int,
)

register_engagement_routes(
    crm_bp,
    actor_name=_actor_name,
    customer_alternate_email_list=_customer_alternate_email_list,
    format_option_dt=_format_option_dt,
    generate_access_key=_generate_access_key,
    get_effective_app_base_url=_get_effective_app_base_url,
    get_meeting_settings=_get_meeting_settings,
    is_valid_email=_is_valid_email,
    is_gathering_request_expired=_is_gathering_request_expired,
    log_opportunity_history=_log_opportunity_history,
    merge_cc_lists=_merge_cc_lists,
    render_system_email_template=_render_system_email_template,
)

register_template_routes(
    crm_bp,
    build_template_macro_values_for_customer=_build_template_macro_values_for_customer,
    email_template_macro_catalog=_email_template_macro_catalog,
    extract_template_macros=_extract_template_macros,
    is_system_template_name=_is_system_template_name,
    is_valid_email=_is_valid_email,
    lead_macro_values=_lead_macro_values,
    lead_sample_columns=_lead_sample_columns,
    log_opportunity_history=_log_opportunity_history,
    sanitize_email_template_body=_sanitize_email_template_body,
    system_template_name_set=_system_template_name_set,
    upsert_lead_from_row=_upsert_lead_from_row,
)

register_lead_routes(
    crm_bp,
    is_valid_email=_is_valid_email,
    log_opportunity_history=_log_opportunity_history,
    upsert_lead_from_row=_upsert_lead_from_row,
)

register_reference_routes(
    crm_bp,
    log_opportunity_history=_log_opportunity_history,
)

register_opportunity_core_routes(
    crm_bp,
    actor_name=_actor_name,
    apply_financial_payload=_apply_financial_payload,
    build_financial_payload=_build_financial_payload,
    build_template_macro_values_for_customer=_build_template_macro_values_for_customer,
    derive_tag_from_action=_derive_tag_from_action,
    ensure_default_cloud_operators=_ensure_default_cloud_operators,
    ensure_default_segments=_ensure_default_segments,
    ensure_default_statuses=_ensure_default_statuses,
    financial_change_lines=_financial_change_lines,
    is_valid_email=_is_valid_email,
    log_opportunity_history=_log_opportunity_history,
    parse_email_list=_parse_email_list,
    status_color_map=_status_color_map,
    tag_color_map=_tag_color_map,
    user_color_map=_user_color_map,
)

register_opportunity_diagram_routes(
    crm_bp,
    actor_name=_actor_name,
    log_history=_log_opportunity_history,
    normalize_macro_key=_normalize_macro_key,
)

register_opportunity_document_routes(
    crm_bp,
    actor_name=_actor_name,
    log_history=_log_opportunity_history,
    sanitize_sow_export_html=_sanitize_sow_export_html,
    generate_sow_docx=_generate_sow_docx,
    generate_html_like_pdf=_generate_html_like_pdf,
    get_sow_brand_config=_get_sow_brand_config,
    next_document_editor_ref_no=_next_document_editor_ref_no,
    compose_master_template_html=_compose_master_template_html,
    render_sow_template_for_customer=_render_sow_template_for_customer,
)

